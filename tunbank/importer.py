"""Controlled balance import (.xlsx / .csv): opening balances and post-reset restoration.

Two steps, exactly as the specification requires:
  1. preview(): read + validate the file and show what WOULD happen. Writes nothing.
  2. commit(): after explicit confirmation, write the batch and the ledger entries in ONE all-or-nothing transaction.

Accepted layouts (header row required):
  WIDE:  nation_id and/or nation_name | money | coal | oil | ... | (loan)     one row per nation
  LONG:  nation_id and/or nation_name | resource | amount                      one row per amount

Nations are identified by nation_id, nation_name, or both (then they must agree). Names are matched
case-insensitively against the live alliance member list; a name that is unknown or shared by two nations blocks
that row - the importer never guesses. Negative resource amounts are real balances (a debt to the alliance) and are
accepted for RESTORE imports. A loan column (loan / outstanding_loan / loan_balance) is stored as an outstanding
loan, never as a deposit.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from . import ledger as L
from . import money as M
from .config import cfg_bool
from .reconcile import batch_fingerprint
from .util import jdump, now_iso, sha256_bytes
from .valuation import value_amounts

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
# A single nation holding more than this in any one resource is treated as impossible.
MAX_PER_RESOURCE = {"money": 10**12 * 100}  # $1 trillion in cents
DEFAULT_MAX_RESOURCE_UNITS = 10**10 * 100    # 10 billion of a resource

ID_HEADERS = ("nation_id", "nationid", "id", "nation_number")
NAME_HEADERS = ("nation_name", "nationname", "name")
ANY_HEADERS = ("nation",)                    # flexible column: a number is an id, anything else is a name
LOAN_HEADERS = ("loan", "loans", "outstanding_loan", "outstanding_loans", "loan_balance", "loan_outstanding")
MODES = ("OPENING", "RESTORE")


def norm_name(text) -> str:
    """Case-insensitive, whitespace-tolerant comparison key for a nation name."""
    return " ".join(str(text or "").split()).casefold()


@dataclass
class Preview:
    filename: str
    file_sha256: str
    mode: str = "OPENING"
    rows: list = field(default_factory=list)          # [(nation_id, resource, units)]  units may be negative (RESTORE)
    nations: dict = field(default_factory=dict)        # nation_id -> {resource: units}
    names: dict = field(default_factory=dict)          # nation_id -> name as PnW knows it
    loans: dict = field(default_factory=dict)          # nation_id -> outstanding loan (cents)
    loan_column: str = ""
    errors: list = field(default_factory=list)         # blocking
    warnings: list = field(default_factory=list)
    missing_members: list = field(default_factory=list)  # alliance members not in the file
    unknown_nations: list = field(default_factory=list)
    totals: dict = field(default_factory=dict)         # NET per resource (signed)
    positive_totals: dict = field(default_factory=dict)
    negative_totals: dict = field(default_factory=dict)  # stored as positive magnitudes
    loan_total_cents: int = 0
    valuation: object | None = None                    # value of the NET totals
    valuation_positive: object | None = None
    valuation_negative: object | None = None
    blocked: bool = False
    content: bytes = b""
    used_names: int = 0                                # rows identified by name only
    used_ids: int = 0


def _read_table(filename: str, content: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raise ValueError("Only .xlsx and .csv files are supported.")


def _cell_units(value, where: str, errors: list, *, allow_negative: bool = False) -> int | None:
    """Strict amount reader. Blank = 0. Returns units or None (error recorded). Never turns a negative into zero."""
    if value is None or str(value).strip() == "":
        return 0
    if isinstance(value, bool):
        errors.append(f"{where}: '{value}' is not a number")
        return None
    text = str(value).strip().replace(",", "").replace("$", "")
    if text.startswith("(") and text.endswith(")"):       # accounting style (1,000) = -1000
        text = "-" + text[1:-1]
    try:
        d = Decimal(text)
    except InvalidOperation:
        errors.append(f"{where}: '{value}' is not a valid number")
        return None
    if not d.is_finite():
        errors.append(f"{where}: '{value}' is not a valid number")
        return None
    if d < 0 and not allow_negative:
        errors.append(f"{where}: negative amount ({value}) - negative balances can only be loaded by a restore "
                      "import after a deposit reset")
        return None
    scaled = d * M.SCALE
    if scaled != scaled.to_integral_value():
        errors.append(f"{where}: more than 2 decimal places ({value})")
        return None
    return int(scaled)


def _as_nation_id(raw):
    """int if the cell is a clean positive whole number, else None."""
    try:
        f = float(str(raw).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None
    if f != int(f) or int(f) <= 0:
        return None
    return int(f)


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def preview(filename: str, content: bytes, *, members: dict | None, snapshot, mode: str = "OPENING") -> Preview:
    if mode not in MODES:
        raise ValueError(f"unknown import mode {mode}")
    p = Preview(filename=filename, file_sha256=sha256_bytes(content), content=content, mode=mode)
    allow_negative = mode == "RESTORE"
    if len(content) > MAX_BYTES:
        p.errors.append("File is larger than 5 MB.")
        p.blocked = True
        return p
    try:
        table = _read_table(filename, content)
    except Exception as exc:  # noqa: BLE001
        p.errors.append(f"Could not read the file: {exc}")
        p.blocked = True
        return p
    table = [r for r in table if any(c not in (None, "") for c in r)]
    if len(table) < 2:
        p.errors.append("The file has no data rows.")
        p.blocked = True
        return p
    if len(table) - 1 > MAX_ROWS:
        p.errors.append(f"Too many rows (max {MAX_ROWS}).")
        p.blocked = True
        return p
    if members is None:
        p.errors.append("Could not read the alliance member list from PnW, so nations cannot be validated. "
                        "Try again in a minute.")
        p.blocked = True
        return p

    # ---- header
    header = [str(h).strip().lower().replace(" ", "_") if h is not None else "" for h in table[0]]
    hmap: dict = {}
    for i, h in enumerate(header):
        if h in ID_HEADERS:
            hmap.setdefault("id", i)
        elif h in NAME_HEADERS:
            hmap.setdefault("name", i)
        elif h in ANY_HEADERS:
            hmap.setdefault("any", i)
        elif h in LOAN_HEADERS:
            if "loan" in hmap:
                p.errors.append("There is more than one loan column.")
            hmap["loan"] = i
            p.loan_column = str(table[0][i]).strip()
        elif h in ("resource", "type"):
            hmap["resource"] = i
        elif h in ("amount", "quantity", "qty", "value"):
            hmap["amount"] = i
        else:
            try:
                res = M.resolve_resource(h)
            except M.AmountError:
                if h:
                    p.warnings.append(f"Column '{table[0][i]}' is not a known resource and was ignored.")
                continue
            if res in hmap:
                p.errors.append(f"Resource '{res}' appears in two columns.")
            hmap[res] = i
    if not ({"id", "name", "any"} & set(hmap)):
        p.errors.append("The file needs a 'nation_id' column, a 'nation_name' column, or both.")
        p.blocked = True
        return p
    long_layout = "resource" in hmap and "amount" in hmap
    wide_cols = [r for r in M.RESOURCES if r in hmap]
    if not long_layout and not wide_cols and "loan" not in hmap:
        p.errors.append("No resource columns found (expected money, coal, oil... or resource+amount).")
        p.blocked = True
        return p

    # ---- name index (live alliance members)
    by_name: dict = {}
    for nid, nm in members.items():
        by_name.setdefault(norm_name(nm), []).append(int(nid))

    seen_nation_rows: dict = {}      # nation_id -> row number (wide layout)
    seen_pairs = set()
    seen_rows = set()
    unknown_ids = set()

    def identify(lineno, row):
        """Return the nation_id for this row, or None after recording a blocking error."""
        id_raw = row[hmap["id"]] if "id" in hmap else None
        name_raw = row[hmap["name"]] if "name" in hmap else None
        any_raw = row[hmap["any"]] if "any" in hmap else None
        if not _blank(any_raw):                      # flexible column: number -> id, text -> name
            if _as_nation_id(any_raw) is not None and _blank(id_raw):
                id_raw = any_raw
            elif _blank(name_raw):
                name_raw = any_raw
        has_id, has_name = not _blank(id_raw), not _blank(name_raw)
        if not has_id and not has_name:
            p.errors.append(f"Row {lineno}: no nation_id or nation_name.")
            return None
        nid = None
        if has_id:
            nid = _as_nation_id(id_raw)
            if nid is None:
                p.errors.append(f"Row {lineno}: '{id_raw}' is not a valid nation id.")
                return None
        if has_name:
            key = norm_name(name_raw)
            hits = by_name.get(key, [])
            if has_id:
                if nid in members and norm_name(members[nid]) != key:
                    p.errors.append(f"Row {lineno}: nation_id {nid} is '{members[nid]}' in PnW but the file says "
                                    f"'{str(name_raw).strip()}' - id and name disagree, row blocked.")
                    return None
                if nid not in members and hits:
                    p.errors.append(f"Row {lineno}: nation_id {nid} does not match the name "
                                    f"'{str(name_raw).strip()}' (that name belongs to nation {hits[0]}) - row blocked.")
                    return None
            else:
                if not hits:
                    p.errors.append(f"Row {lineno}: nation name '{str(name_raw).strip()}' was not found in the "
                                    "alliance - row blocked, nothing guessed.")
                    return None
                if len(hits) > 1:
                    p.errors.append(f"Row {lineno}: nation name '{str(name_raw).strip()}' is ambiguous "
                                    f"(nations {', '.join(map(str, sorted(hits)))}) - row blocked.")
                    return None
                nid = hits[0]
        if has_id:
            p.used_ids += 1
        else:
            p.used_names += 1
        if nid not in members:
            unknown_ids.add(nid)
            return None          # reported once, in aggregate, below
        p.names[nid] = members[nid]
        return nid

    for lineno, row in enumerate(table[1:], start=2):
        row = list(row) + [None] * (len(header) - len(row))
        key = tuple("" if c is None else str(c).strip() for c in row)
        if key in seen_rows:
            p.errors.append(f"Row {lineno}: exact duplicate of an earlier row.")
            continue
        seen_rows.add(key)
        nid = identify(lineno, row)
        if nid is None:
            continue
        entries = []
        loan_cell = None
        if long_layout:
            rname = row[hmap["resource"]]
            if str(rname).strip().lower().replace(" ", "_") in LOAN_HEADERS:
                loan_cell = row[hmap["amount"]]
                p.loan_column = p.loan_column or "loan"
            else:
                try:
                    res = M.resolve_resource(str(rname))
                except M.AmountError:
                    p.errors.append(f"Row {lineno}: unknown resource '{rname}'.")
                    continue
                u = _cell_units(row[hmap["amount"]], f"Row {lineno} {res}", p.errors, allow_negative=allow_negative)
                if u:
                    entries.append((res, u))
        else:
            if nid in seen_nation_rows:
                p.errors.append(f"Row {lineno}: nation {nid} ({members[nid]}) is a duplicate - it already appears "
                                f"in row {seen_nation_rows[nid]}.")
                continue
            seen_nation_rows[nid] = lineno
            for res in wide_cols:
                u = _cell_units(row[hmap[res]], f"Row {lineno} {res}", p.errors, allow_negative=allow_negative)
                if u:
                    entries.append((res, u))
            if "loan" in hmap:
                loan_cell = row[hmap["loan"]]
        p.nations.setdefault(nid, {})
        for res, u in entries:
            if (nid, res) in seen_pairs:
                p.errors.append(f"Row {lineno}: nation {nid} already has {res} in this file (duplicate).")
                continue
            seen_pairs.add((nid, res))
            cap = MAX_PER_RESOURCE.get(res, DEFAULT_MAX_RESOURCE_UNITS)
            if abs(u) > cap:
                p.errors.append(f"Row {lineno}: {res} amount for nation {nid} is impossibly large.")
                continue
            p.rows.append((nid, res, u))
            p.nations[nid][res] = u
        if not _blank(loan_cell):
            lu = _cell_units(loan_cell, f"Row {lineno} loan", p.errors, allow_negative=False)
            if lu:
                if nid in p.loans:
                    p.errors.append(f"Row {lineno}: nation {nid} already has a loan in this file (duplicate).")
                elif lu > MAX_PER_RESOURCE["money"]:
                    p.errors.append(f"Row {lineno}: loan for nation {nid} is impossibly large.")
                else:
                    p.loans[nid] = lu

    if unknown_ids:
        ids = sorted(unknown_ids)
        p.unknown_nations = ids
        p.errors.append(f"{len(ids)} nation(s) are not in the alliance: "
                        + ", ".join(map(str, ids[:15])) + (" ..." if len(ids) > 15 else ""))
    if not p.rows and not p.loans and not p.errors:
        p.errors.append("The file contains no amounts.")

    for _, res, u in p.rows:
        p.totals[res] = p.totals.get(res, 0) + u
        if u > 0:
            p.positive_totals[res] = p.positive_totals.get(res, 0) + u
        else:
            p.negative_totals[res] = p.negative_totals.get(res, 0) - u
    p.loan_total_cents = sum(p.loans.values())
    p.valuation = value_amounts(p.totals, snapshot)
    p.valuation_positive = value_amounts(p.positive_totals, snapshot)
    p.valuation_negative = value_amounts(p.negative_totals, snapshot)
    p.missing_members = sorted(n for n in members if n not in p.nations)
    if p.missing_members:
        p.warnings.append(f"{len(p.missing_members)} current alliance member(s) are not in the file "
                          "(they will start with zero).")
    if not p.valuation.complete:
        p.warnings.append("Some resource prices are missing: the market value is incomplete.")
    if p.valuation.stale or p.valuation.suspicious:
        p.warnings.append("Market prices are stale/unusual: treat the value as approximate.")
    if p.loans:
        p.warnings.append("Loan amounts are recorded as OUTSTANDING LOANS, not deposits. They stay on file until "
                          "the loan module is built.")
    p.blocked = bool(p.errors)
    return p


def commit(conn, p: Preview, *, admin_id: str, note: str, snapshot_id, kind: str = "OPENING",
           reset_id: int | None = None) -> dict:
    """Write the batch. Must be called inside db.tx(). Refuses anything blocked.

    kind='OPENING': first-time balances (positive only, one per nation/resource, ever).
    kind='RESTORE': balances loaded after a /deposit reset; may include negatives; tied to that reset."""
    if p.blocked or p.errors:
        raise L.LedgerError("This import has errors and cannot be committed.")
    if kind not in MODES or p.mode != kind:
        raise L.LedgerError("This preview was made for a different kind of import.")
    L.assert_can_mutate(conn, None, "import")
    if kind == "OPENING":
        if not cfg_bool(conn, "opening_import_allowed"):
            raise L.LedgerError("Opening-balance import is switched off.")
        if any(u < 0 for _, _, u in p.rows):
            raise L.LedgerError("Negative balances can only be loaded by a restore import after a deposit reset.")
        dup = conn.execute("SELECT id FROM import_batches WHERE source_sha256=?", (p.file_sha256,)).fetchone()
        if dup:
            raise L.LedgerError(f"This exact file was already imported as batch #{dup['id']}.")
        already = conn.execute(
            "SELECT nation_id, resource FROM ledger_entries WHERE entry_type IN ('OPENING','RESTORE')").fetchall()
        have = {(r["nation_id"], r["resource"]) for r in already}
        clash = [f"{n}/{r}" for n, r, _ in p.rows if (n, r) in have]
        if clash:
            raise L.LedgerError("These nation/resource pairs already have an imported balance: "
                                + ", ".join(clash[:10]) + ". Opening balances can never be overwritten; "
                                "use a documented /bank adjust instead (or a /deposit reset + restore).")
    else:
        if reset_id is None:
            raise L.LedgerError("A restore import must be tied to a deposit reset.")
        rs = conn.execute("SELECT id, restore_batch_id FROM deposit_resets WHERE id=?", (reset_id,)).fetchone()
        if not rs:
            raise L.LedgerError(f"Deposit reset #{reset_id} does not exist.")
        if rs["restore_batch_id"] is not None:
            raise L.LedgerError(f"Deposit reset #{reset_id} was already restored by batch #{rs['restore_batch_id']}.")
    fingerprint_rows = [{"nation_id": n, "resource": r, "amount": u} for n, r, u in p.rows]
    totals = {"net": p.totals, "positive": p.positive_totals, "negative": p.negative_totals,
              "loans_cents": p.loan_total_cents, "loan_nations": len(p.loans)}
    cur = conn.execute(
        "INSERT INTO import_batches(created_at,admin_discord_id,source_filename,source_sha256,row_count,"
        "totals_json,rows_sha256,price_snapshot_id,value_cents,note,kind,reset_id) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
        (now_iso(), str(admin_id), p.filename, p.file_sha256, len(p.rows), jdump(totals),
         batch_fingerprint(fingerprint_rows), snapshot_id, p.valuation.total_cents if p.valuation else None,
         note, kind, reset_id))
    batch_id = cur.lastrowid
    conn.execute("INSERT INTO import_files(batch_id, content) VALUES(?,?)", (batch_id, p.content))
    for n, r, u in p.rows:
        conn.execute("INSERT INTO opening_balance_rows(batch_id,nation_id,resource,amount) VALUES(?,?,?,?)",
                     (batch_id, n, r, u))
    for n in p.nations:
        L.ensure_member(conn, n, p.names.get(n))
    g = L.new_group("OPEN" if kind == "OPENING" else "REST")
    L.post_entries(conn, [
        dict(group_id=g, nation_id=n, bucket="AVAILABLE", resource=r, delta=u, entry_type=kind,
             batch_id=batch_id, actor=admin_id,
             note=(f"Opening balance import #{batch_id}" if kind == "OPENING"
                   else f"Restore import #{batch_id} after deposit reset #{reset_id}"),
             price_snapshot_id=snapshot_id)
        for n, r, u in p.rows])
    for n, cents in sorted(p.loans.items()):
        conn.execute("INSERT INTO imported_loans(batch_id,nation_id,outstanding_cents,source_column,created_at) "
                     "VALUES(?,?,?,?,?)", (batch_id, n, cents, p.loan_column or "loan", now_iso()))
    if kind == "RESTORE":
        conn.execute("UPDATE deposit_resets SET restore_batch_id=? WHERE id=? AND restore_batch_id IS NULL",
                     (batch_id, reset_id))
    from . import configaudit as CA
    CA.record(conn, actor=admin_id,
              setting="opening_balance_import" if kind == "OPENING" else "deposit_restore_import",
              previous="—",
              new=f"batch #{batch_id}: {len(p.rows)} amounts for {len(p.nations)} nations",
              target=f"{p.filename} · {note}" + (f" · reset #{reset_id}" if reset_id else ""),
              category="IMPORT", only_if_changed=False)
    L.audit(conn, admin_id, "OPENING_IMPORT" if kind == "OPENING" else "RESTORE_IMPORT", f"batch:{batch_id}", {
        "filename": p.filename, "sha256": p.file_sha256, "rows": len(p.rows), "totals": p.totals,
        "negative_totals": p.negative_totals, "reset_id": reset_id, "note": note})
    if p.loans:
        CA.record(conn, actor=admin_id, setting="loan_import", previous="—",
                  new=f"batch #{batch_id}: {len(p.loans)} outstanding loans, total {p.loan_total_cents / 100:,.2f}",
                  target=f"{p.filename} · stored as outstanding loans, not deposits",
                  category="IMPORT", only_if_changed=False)
        L.audit(conn, admin_id, "LOAN_IMPORT", f"batch:{batch_id}", {
            "loans": {str(n): c for n, c in sorted(p.loans.items())}, "total_cents": p.loan_total_cents,
            "column": p.loan_column})
    return {"batch_id": batch_id, "rows": len(p.rows), "nations": len(p.nations), "totals": p.totals,
            "loans": len(p.loans), "loan_total_cents": p.loan_total_cents, "kind": kind}
