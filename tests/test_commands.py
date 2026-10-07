"""Drives the real command code with a fake Discord. Run: python -m unittest discover -s tests -v"""
import asyncio
import re
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import fake_discord  # noqa: E402

fake_discord.install()

import discord  # noqa: E402
from discord import app_commands  # noqa: E402

from test_core import ALLIANCE, PRICES, FakePnW, rec  # noqa: E402
from tunbank import cmds_admin, cmds_audit, cmds_backup, cmds_bulk, cmds_chart, cmds_econ, cmds_grant, cmds_help, cmds_market, cmds_member, cmds_offshore  # noqa: E402
from tunbank import ledger as L  # noqa: E402
from tunbank import reconcile as R  # noqa: E402
from tunbank import money as M  # noqa: E402
from tunbank.alerts import AlertService  # noqa: E402
from tunbank.config import BankAccess, Settings  # noqa: E402
from tunbank import credentials as CR  # noqa: E402
from tunbank.memberdeposit import MemberDepositService  # noqa: E402
from tunbank.offshore import OffshoreService  # noqa: E402
from tunbank.db import Database  # noqa: E402
from tunbank.pnw import PnWRejected  # noqa: E402
from tunbank.scanner import Scanner  # noqa: E402
from tunbank.ui import Services  # noqa: E402
from tunbank.valuation import PriceService  # noqa: E402
from tunbank.withdrawals import WithdrawalService  # noqa: E402

FI = fake_discord.FakeInteraction


class PnW2(FakePnW):
    async def fetch_nation(self, nid):
        return {"id": nid, "nation_name": f"Nation{nid}", "alliance_id": ALLIANCE if nid != 55 else 7,
                "alliance_position": "MEMBER", "discord": {1: "alice", 2: "bob"}.get(nid, "")}


class Cmd(unittest.TestCase):
    def make_settings(self):
        return Settings("t", "k", "b", "k", ALLIANCE, Path(self.tmp.name), {900001})

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "t.db")
        self.db.migrate()
        self.pnw = PnW2()
        self.settings = self.make_settings()
        prices = PriceService(self.db, self.pnw)
        self.svc = Services(self.settings, self.db, self.pnw, prices, Scanner(self.db, self.pnw, prices, self.settings),
                            WithdrawalService(self.db, self.pnw, prices, self.settings), AlertService(self.db))
        self.svc.offshore = OffshoreService(self.db, self.pnw, prices, self.settings)
        self.svc.crypto = CR.Crypto(self.settings.credential_key)
        self.svc.deposits = MemberDepositService(self.db, self.pnw, prices, self.settings, self.svc.crypto, self.svc.scanner)
        self.bank = app_commands.Group("bank")
        self.bankset = app_commands.Group("bankset")
        self.nation = app_commands.Group("nation")
        self.ledger = app_commands.Group("ledger")
        self.tax = app_commands.Group("tax")
        self.audit = app_commands.Group("audit")
        self.chart = app_commands.Group("chart")
        self.bulk = app_commands.Group("bulk")
        self.grant = app_commands.Group("grant")
        cmds_member.register(self.bank, self.nation, self.svc)
        cmds_econ.register(self.bank, self.svc)
        cmds_admin.register(self.bank, self.bankset, self.ledger, self.svc)
        cmds_backup.register(self.bankset, self.svc)
        cmds_chart.register(self.chart, self.svc)
        cmds_bulk.register(self.bulk, self.svc)
        cmds_grant.register(self.grant, self.svc)
        cmds_offshore.register(self.bank, self.svc)
        cmds_audit.register_tax(self.tax, self.svc)
        cmds_audit.register_audit(self.audit, self.ledger, self.svc)
        from discord.ext import commands as _c
        self.tree = _c.Bot().tree
        cmds_help.register(self.tree, self.svc)
        cmds_market.register(self.tree, self.svc)
        self.alice = discord.User(111, "alice")
        self.admin = discord.User(900001, "owner")      # owner => ADMIN
        self.econ2 = discord.User(222, "second")

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def run_async(self, c):
        return asyncio.run(c)

    def call(self, group, name, user, *a, auto_confirm=True, **kw):
        i = FI(user, auto_confirm=auto_confirm)
        i.command = __import__("types").SimpleNamespace(qualified_name=f"{group.name} {name}")
        self.run_async(getattr(group, "commands")[name].callback(i, *a, **kw))
        return i

    def fund_alice(self, amounts):
        """Real path: baseline scan, link, real deposit detected by scanner."""
        self.run_async(self.svc.scanner.scan())
        self.call(self.nation, "link", self.alice, 1)
        self.pnw.recs.append(rec(900, 1, amounts))
        self.pnw.holdings = {k: int(v * 100) for k, v in amounts.items()}  # bank really holds it
        self.run_async(self.svc.scanner.scan())

    def test_command_registration_is_valid(self):
        names = sorted(self.bank.commands)
        self.assertIn("withdrawself", names)
        self.assertIn("settransferlimit", self.bankset.commands)
        for g in (self.bank, self.bankset, self.nation, self.tax, self.audit, self.ledger):
            self.assertLessEqual(len(g.commands), 25, g.name)

    def test_link_requires_matching_discord_name_and_alliance(self):
        i = self.call(self.nation, "link", self.alice, 1)
        self.assertIn("Linked", i.text())
        i = self.call(self.nation, "link", discord.User(333, "mallory"), 2)      # nation 2 is bob's
        self.assertIn("can't verify", i.text())
        i = self.call(self.nation, "link", discord.User(333, "mallory"), 55)     # other alliance
        self.assertIn("not in our alliance", i.text())

    def test_dashboard_shows_quantity_and_value_but_never_tax(self):
        self.fund_alice({"money": 1000000, "aluminum": 1000, "coal": 1000})
        self.pnw.recs.append(rec(901, 1, {"money": 5555}, "tax", tax_id=3))
        self.run_async(self.svc.scanner.scan())
        i = self.call(self.bank, "dashboard", self.alice)
        t = i.text()
        self.assertIn("Current Market Value", t)
        self.assertIn("$1,000,000.00", t)
        self.assertNotIn("5,555", t)
        self.assertNotIn("tax", t.lower().replace("taxes", ""))

    def test_withdrawself_needs_confirmation_then_sends(self):
        self.fund_alice({"money": 1000000})
        i = self.call(self.bank, "withdrawself", self.alice, "money=400k", auto_confirm=False)
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.assertIn("Cancelled", i.text())
        i = self.call(self.bank, "withdrawself", self.alice, "money=400k")
        self.assertEqual(self.pnw.withdraw_calls, 1)
        self.assertIn("Withdrawal completed", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 60000000)

    def test_withdrawself_locked_funds_rejected(self):
        self.fund_alice({"money": 1000000})
        self.call(self.bank, "reserve", self.admin, 1, "money=1m", "WARCHEST", "war prep")
        i = self.call(self.bank, "withdrawself", self.alice, "money=1")
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.assertIn("Locked", i.text())

    def test_reserve_release_no_pnw_call_and_shows_value(self):
        self.fund_alice({"money": 1000000, "coal": 100})
        i = self.call(self.bank, "reserve", self.admin, 1, "money=500k coal=50", "WARCHEST", "war")
        self.assertIn("Current Market Value", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "LOCKED")["money"], 50000000)
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 50000000)
        i = self.call(self.bank, "release", self.admin, 1, "done")
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "LOCKED"), {})
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000)

    def test_non_staff_cannot_use_staff_commands(self):
        for name, args in (("reserve", (1, "money=1", "X", "r")), ("holdings", ()), ("adjust", (1, "r", "e")),
                           ("importopening", (discord.Attachment("a.csv", b"x"), "n"))):
            i = self.call(self.bankset if name == "importopening" else self.bank, name, self.alice, *args)
            self.assertIn("permission", i.text(), name)
        for name, args in (("setrole", (discord.app_commands.Choice("ADMIN", "ADMIN"), discord.Role(1))),
                           ("backup", ()), ("config", ())):
            i = self.call(self.bankset, name, self.alice, *args)
            self.assertIn("permission", i.text(), name)

    def test_adjust_positive_needs_second_person(self):
        self.fund_alice({"money": 1000})
        i = self.call(self.bank, "adjust", self.admin, 1, "missed deposit", "ticket 9", add="money=100")
        self.assertIn("Approval request", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000)
        # same person can't approve
        i = self.call(self.bank, "approve", self.admin, 1)
        self.assertIn("different staff member", i.text())
        # second person (give them Minister via the owner-only setrole path)
        with self.db.tx() as c:
            c.execute("INSERT INTO bankers(discord_id,added_by,added_at) VALUES('222','x','x')")
            c.execute("INSERT INTO role_permissions(level,role_id) VALUES('MINISTER','77')")
        self.econ2.roles = [discord.Role(77)]
        i = self.call(self.bank, "approve", self.econ2, 1)
        self.assertIn("Approved", i.text())
        i = self.call(self.bank, "adjust", self.admin, 1, "missed deposit", "ticket 9", add="money=100")
        self.assertIn("Accounting adjustment", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 110000)

    def test_large_econ_withdrawal_needs_approval(self):
        self.fund_alice({"money": 100000000})
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "approval_threshold_value", "1000", "t")
            c.execute("INSERT INTO role_permissions(level,role_id) VALUES('MINISTER','77')")
        self.econ2.roles = [discord.Role(77)]
        ch = discord.app_commands.Choice("ALLIANCE", "MEMBER_AVAILABLE")
        i = self.call(self.bank, "withdraw", self.admin, ch, 1, "money=5m", "payout", member="1")
        self.assertIn("second approver", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.call(self.bank, "approve", self.econ2, 1)
        i = self.call(self.bank, "withdraw", self.admin, ch, 1, "money=5m", "payout", member="1")
        self.assertEqual(self.pnw.withdraw_calls, 1)
        self.assertIn("completed", i.text())

    def test_import_preview_blocks_bad_file_and_good_file_commits(self):
        bad = discord.Attachment("b.csv", b"nation_id,money\n1,-5\n99,3\n")
        i = self.call(self.bankset, "importopening", self.admin, bad, "from locutus")
        self.assertIn("BLOCKED", i.text())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0], 0)
        good = discord.Attachment("g.csv", b"nation_id,money,coal\n1,500,20\n2,100,0\n")
        i = self.call(self.bankset, "importopening", self.admin, good, "from locutus", auto_confirm=False)
        self.assertIn("cancelled", i.text().lower())
        i = self.call(self.bankset, "importopening", self.admin, good, "from locutus")
        self.assertIn("imported", i.text().lower())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE"), {"money": 50000, "coal": 2000})

    def test_vault_shows_alliance_owned_and_tax(self):
        self.fund_alice({"money": 1000000})
        self.pnw.recs.append(rec(902, 2, {"money": 777}, "tax", tax_id=3))
        self.pnw.recs.append(rec(903, 2, {"money": 50}, "#ignore"))
        self.pnw.holdings = {"money": 100000000 + 300000 + 77700 + 5000}
        self.run_async(self.svc.scanner.scan())
        i = self.call(self.bank, "holdings", self.admin)
        t = i.text()
        for word in ("Real PnW bank", "ALLIANCE-OWNED", "Taxes collected", "#ignore", "$777.00"):
            self.assertIn(word, t)

    def test_exports_and_tax_commands(self):
        self.fund_alice({"money": 1000000})
        self.pnw.recs.append(rec(905, 2, {"money": 777}, "tax", tax_id=3))
        self.run_async(self.svc.scanner.scan())
        for kind in ("balances", "ledger", "deposits", "withdrawals", "locks", "tax", "audit", "integrity", "vault"):
            i = self.call(self.bank, "records", self.admin, discord.app_commands.Choice(kind, kind))
            files = [m.get("file") for m in i.sent if m.get("file")]
            self.assertEqual(len(files), 1, kind)
            self.assertGreater(len(files[0].fp.getvalue()), 1000)
        i = self.call(self.tax, "dashboard", self.admin)
        self.assertIn("$777.00", i.text())
        i = self.call(self.tax, "export", self.admin)
        self.assertTrue(any(m.get("file") for m in i.sent))

    def test_reconcile_and_emergency_lock_commands(self):
        self.fund_alice({"money": 1000000})
        i = self.call(self.bank, "reconcile", self.admin)
        self.assertIn("Reconciliation", i.text())
        ch = discord.app_commands.Choice("on", "on")
        self.call(self.ledger, "emergencylock", self.admin, ch, "test drill")
        i = self.call(self.bank, "withdrawself", self.alice, "money=1")
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.assertIn("EMERGENCY LOCK", i.text())
        i = self.call(self.ledger, "dashboard", self.admin)
        self.assertIn("EMERGENCY_LOCK", i.text())

    def test_hash_chains_stay_valid_after_real_command_use(self):
        """Regression: Discord IDs are numbers; the tamper-check must still verify."""
        from tunbank import reconcile as R
        self.fund_alice({"money": 1000000})
        self.call(self.bank, "reserve", self.admin, 1, "money=1k", "X", "r")
        self.call(self.bank, "withdrawself", self.alice, "money=10")
        self.call(self.bank, "freeze", self.admin, 1, "test")
        self.call(self.bankset, "setrole", self.admin, discord.app_commands.Choice("BANKER", "BANKER"), discord.Role(5))
        with self.db.read() as c:
            self.assertEqual(R.check_chains(c), [])
        i = self.call(self.bank, "reconcile", self.admin)
        self.assertIn("Reconciliation OK", i.text(), i.text())

    def test_nation_fields_accept_id_name_link_username_and_mention(self):
        self.fund_alice({"money": 1000})
        for typed in ("1", "#1", "ALPHA", "alpha", "alice", "@Alice", "<@111>", "Alpha",
                      "https://politicsandwar.com/nation/id=1"):
            i = self.call(self.bank, "freeze", self.admin, typed, "test")
            self.assertIn("Account #1 frozen", i.text(), typed)
        self.call(self.bank, "unfreeze", self.admin, "alice", "done")
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT frozen FROM members WHERE nation_id=1").fetchone()[0], 0)

    def test_nation_lookup_never_guesses(self):
        self.fund_alice({"money": 1000})
        i = self.call(self.bank, "freeze", self.admin, "Alph", "x")      # only part of a name
        self.assertIn("couldn't find", i.text())
        self.assertIn("Did you mean", i.text())
        self.assertIn("[#1]", i.text())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT frozen FROM members WHERE nation_id=1").fetchone()[0], 0)
        i = self.call(self.bank, "freeze", self.admin, "nobody-here", "x")
        self.assertIn("couldn't find", i.text())
        with self.db.tx() as c:
            L.ensure_member(c, 2, "alpha")                                  # two nations, same name
        i = self.call(self.bank, "freeze", self.admin, "alpha", "x")
        self.assertIn("More than one nation", i.text())
        i = self.call(self.bank, "freeze", self.admin, "<@999>", "x")        # nobody linked
        self.assertIn("has not linked", i.text())

    def test_nation_autocomplete_suggests_names(self):
        self.fund_alice({"money": 1000})
        cb = self.bank.commands["reserve"].autocompletes["nation"]
        choices = self.run_async(cb(FI(self.admin), "alp"))
        self.assertTrue(any(c.value == "1" and "Alpha" in c.name for c in choices))
        self.assertIn("withdraw", self.bank.commands)
        self.assertEqual(set(self.bank.commands["withdraw"].autocompletes), {"destination", "member"})
        self.assertEqual(self.run_async(cb(FI(self.admin), "zzzz")), [])

    def test_icons_are_configurable_without_code_changes(self):
        from tunbank import icons
        try:
            self.fund_alice({"money": 1000000, "oil": 100})
            i = self.call(self.bank, "dashboard", self.alice)
            self.assertIn("💵", i.text())
            self.assertIn("🛢️", i.text())
            money = discord.app_commands.Choice("Cash", "money")
            i = self.call(self.bankset, "seticon", self.alice, money, "<:cash:123456789012345678>")
            self.assertIn("permission", i.text())
            i = self.call(self.bankset, "seticon", self.admin, money, "abc")
            self.assertIn("doesn't look like an emoji", i.text())
            self.call(self.bankset, "seticon", self.admin, money, "<:cash:123456789012345678>")
            i = self.call(self.bank, "dashboard", self.alice)
            self.assertIn("<:cash:123456789012345678>", i.text())
            self.call(self.bankset, "seticon", self.admin, money, "default")
            i = self.call(self.bank, "dashboard", self.alice)
            self.assertIn("💵", i.text())
            i = self.call(self.bankset, "icons", self.admin)
            self.assertIn("Resource icons", i.text())
        finally:
            icons.load({})

    def test_long_lists_are_paginated_and_short_ones_are_not(self):
        self.fund_alice({"money": 1000})
        with self.db.tx() as c:
            for n in range(9):
                c.execute("INSERT INTO transactions(created_at,updated_at,tx_type,status,funding_source,member_nation_id,"
                          "dest_nation_id,actor_discord_id,idempotency_key,value_cents) VALUES('2026-10-01T00:00:00Z','x',"
                          "'WITHDRAW_SELF','FAILED','MEMBER_AVAILABLE',1,1,'111',?,500)", (f"k{n}",))
        i = self.call(self.audit, "transactions", self.admin)
        self.assertIsNotNone(i.sent[-1].get("view"))
        self.assertIn("Page 1 of 2", i.sent[-1]["embed"].footer_text if hasattr(i.sent[-1]["embed"], "footer_text") else "Page 1 of 2")
        i = self.call(self.bank, "transactions", self.admin, status="NOPE")
        self.assertIsNone(i.sent[-1].get("view"))

    def test_staff_lists_and_vault_render(self):
        self.fund_alice({"money": 1000})
        self.pnw.recs.append(rec(950, 77, {"money": 5}))     # a non-member sender -> needs a human decision
        self.run_async(self.svc.scanner.scan())
        i = self.call(self.bank, "review", self.admin)
        self.assertIn("needs a decision", i.text())
        self.assertIn("PnW record #950", i.text())
        view = i.view()
        self.assertEqual(list(view.handlers), ["Credit to a nation…", "Alliance money", "Dismiss"])
        pressed = self.press(view, "Credit to a nation…", self.admin)
        pressed.modal.children[0].value = "alice"
        pressed.modal.children[1].value = "member paid with a typo in the note"
        sub = FI(self.admin)
        self.run_async(pressed.modal.on_submit(sub))
        self.assertIn("resolved (credit)", sub.text())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000 + 500)
        for grp, name, args in ((self.bank, "approvals", ()), (self.audit, "stafflog", ()), (self.bank, "holdings", ()),
                                (self.bank, "history", ()), (self.audit, "nation", ("1",)), (self.tax, "report", ())):
            user = self.alice if name == "history" else self.admin
            i = self.call(grp, name, user, *args)
            self.assertTrue(i.sent, name)

    def test_charts_for_members_and_staff(self):
        self.fund_alice({"money": 1000000, "coal": 500})
        self.pnw.recs.append(rec(960, 2, {"money": 777}, "tax", tax_id=3))
        self.run_async(self.svc.scanner.scan())
        self.call(self.bank, "reconcile", self.admin)
        for name in ("mybalance", "mytrend"):
            i = self.call(self.chart, name, self.alice)
            files = [m.get("file") for m in i.sent if m.get("file")]
            self.assertEqual(len(files), 1, name)
            self.assertTrue(files[0].fp.getvalue().startswith(b"\x89PNG"), name)
            self.assertEqual(i.sent[-1]["embed"].image_url, "attachment://" + files[0].filename)
        for name, args in (("vault", ()), ("members", ()), ("deposits", ()), ("tax", ()),
                           ("nation", ("alice", discord.app_commands.Choice("Resource mix", "mix"))),
                           ("nation", ("1", discord.app_commands.Choice("Value over time", "trend"))),
                           ("nation", ("1", discord.app_commands.Choice("Deposits", "deposits")))):
            i = self.call(self.chart, name, self.admin, *args)
            self.assertTrue(any(m.get("file") for m in i.sent), (name, i.text()))
        # members must not see staff charts (tax especially)
        for name in ("vault", "tax", "deposits", "members"):
            i = self.call(self.chart, name, self.alice)
            self.assertIn("permission", i.text(), name)
        # no data -> a friendly message, not a crash
        i = self.call(self.chart, "mytrend", discord.User(555, "stranger"))
        self.assertIn("Link your nation", i.text())

    def press(self, view, label, user, auto_confirm=True):
        i = FI(user, auto_confirm=auto_confirm)
        self.run_async(view.handlers[label](i))
        return i

    def test_member_buttons_withdraw_form_and_shortcuts(self):
        self.fund_alice({"money": 1000000, "coal": 500})
        for cmd in ("dashboard", "balance"):
            view = self.call(self.bank, cmd, self.alice).view()
            self.assertEqual(list(view.handlers), ["Withdraw", "Withdraw all cash", "History", "Chart", "How to deposit", "Refresh"])
        view = self.call(self.bank, "dashboard", self.alice).view()
        # someone else cannot press my buttons
        self.assertIn("belongs to someone else", self.press(view, "Withdraw", self.admin).text())
        # Withdraw -> pop-up form -> confirmation -> real withdrawal
        i = self.press(view, "Withdraw", self.alice)
        self.assertEqual([c.label for c in i.modal.children], ["What to withdraw", "Note (optional)"])
        i.modal.children[0].value = "money=100k"
        sub = FI(self.alice)
        self.run_async(i.modal.on_submit(sub))
        self.assertEqual(self.pnw.withdraw_calls, 1)
        self.assertIn("Withdrawal completed", sub.text())
        self.assertEqual(list(sub.view().handlers), ["My dashboard", "Withdraw more"])
        # bad text in the form is explained, nothing sent
        i = self.press(view, "Withdraw", self.alice)
        i.modal.children[0].value = "lots of money"
        sub = FI(self.alice)
        self.run_async(i.modal.on_submit(sub))
        self.assertIn("couldn't read", sub.text())
        self.assertEqual(self.pnw.withdraw_calls, 1)
        # cancelling on the confirmation screen sends nothing
        i = self.press(view, "Withdraw", self.alice)
        i.modal.children[0].value = "money=1k"
        sub = FI(self.alice, auto_confirm=False)
        self.run_async(i.modal.on_submit(sub))
        self.assertEqual(self.pnw.withdraw_calls, 1)
        # one-click: all available cash
        i = self.press(view, "Withdraw all cash", self.alice)
        self.assertEqual(self.pnw.withdraw_calls, 2)
        with self.db.read() as c:
            self.assertNotIn("money", L.get_balances(c, 1, "AVAILABLE"))
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["coal"], 50000)
        self.assertIn("no available cash", self.press(view, "Withdraw all cash", self.alice).text())

    def test_member_buttons_history_chart_refresh_and_deposit_check(self):
        self.fund_alice({"money": 1000000, "coal": 500})
        view = self.call(self.bank, "dashboard", self.alice).view()
        self.assertIn("Your history", self.press(view, "History", self.alice).text())
        i = self.press(view, "Chart", self.alice)
        self.assertTrue(any(m.get("file") for m in i.sent))
        i = self.press(view, "Refresh", self.alice)
        self.assertIn("Your TUN Bank account", i.edited["embed"].title)
        i = self.press(view, "How to deposit", self.alice)
        self.assertIn("How to deposit", i.text())
        dep_view = i.view()
        self.pnw.recs.append(rec(970, 1, {"money": 5000}))
        i = self.press(dep_view, "Check my deposit now", self.alice)
        self.assertIn("Found **1** new deposit", i.text())
        i = self.press(dep_view, "Check my deposit now", self.alice)       # anti-spam cooldown
        self.assertIn("Please wait", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000 + 500000)

    def test_prices_command_shows_icons_and_real_errors(self):
        i = FI(self.alice)
        self.run_async(self.tree.top["prices"].callback(i))
        t = i.text()
        for word in ("Market prices", "Oil", "🛢️", "Cash", "$100"):
            self.assertIn(word, t)
        from tunbank.pnw import PnWRejected

        async def boom():
            raise PnWRejected('Cannot query field "zzz" on type "Tradeprice".')
        self.pnw.fetch_prices = boom
        view = i.view()
        again = self.press(view, "Refresh", self.alice)
        self.assertIn("Cannot query field", again.edited["embed"].fields[-1][1])
        self.assertIn("Last refresh failed", again.edited["embed"].fields[-1][0])

    def test_staff_buttons_vault_nation_ledger_and_help(self):
        self.fund_alice({"money": 1000000, "coal": 500})
        # vault
        view = self.call(self.bank, "holdings", self.admin).view()
        self.assertEqual(list(view.handlers), ["Chart", "Reconcile now", "Review queue", "Export balances", "Refresh"])
        self.assertIn("Reconciliation", self.press(view, "Reconcile now", self.admin).text())
        self.assertTrue(any(m.get("file") for m in self.press(view, "Export balances", self.admin).sent))
        self.assertTrue(any(m.get("file") for m in self.press(view, "Chart", self.admin).sent))
        self.assertIn("Nothing is waiting", self.press(view, "Review queue", self.admin).text())
        # one nation: charts, reserve via form, freeze via form
        view = self.call(self.audit, "nation", self.admin, "alice").view()
        self.assertEqual(list(view.handlers), ["Resource mix", "Value over time", "Reserve funds…", "Freeze"])
        self.assertTrue(any(m.get("file") for m in self.press(view, "Resource mix", self.admin).sent))
        i = self.press(view, "Reserve funds…", self.admin)
        i.modal.children[0].value, i.modal.children[1].value, i.modal.children[2].value = "money=200k", "WARCHEST", "war prep"
        self.run_async(i.modal.on_submit(FI(self.admin)))
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "LOCKED")["money"], 20000000)
        i = self.press(view, "Freeze", self.admin)
        i.modal.children[0].value = "investigating"
        self.run_async(i.modal.on_submit(FI(self.admin)))
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT frozen FROM members WHERE nation_id=1").fetchone()[0], 1)
        self.assertIn("Unfreeze", list(self.call(self.audit, "nation", self.admin, "alice").view().handlers))
        # a non-staff user opening the staff pop-up gets refused
        view_a = self.call(self.audit, "nation", self.admin, "alice").view()
        self.assertIn("belongs to someone else", self.press(view_a, "Unfreeze", self.alice).text())
        # emergency lock from the integrity dashboard
        view = self.call(self.ledger, "dashboard", self.admin).view()
        i = self.press(view, "Emergency lock", self.admin)
        i.modal.children[0].value = "drill"
        self.run_async(i.modal.on_submit(FI(self.admin)))
        with self.db.read() as c:
            self.assertTrue(L.integrity_state(c)["emergency_lock"])
        self.assertIn("Lift lock", list(self.call(self.ledger, "dashboard", self.admin).view().handlers))
        # help: quick buttons open the same screens
        view = FI(self.admin)
        self.run_async(self.tree.top["help"].callback(view))
        hv = view.view()
        for label in ("My dashboard", "Market prices", "Vault", "Integrity"):
            self.assertIn(label, hv.handlers)
        self.assertIn("Market prices", self.press(hv, "Market prices", self.admin).text())

    def test_approvals_one_per_page_with_approve_button(self):
        self.fund_alice({"money": 1000})
        self.call(self.bank, "adjust", self.admin, "alice", "missed deposit", "ticket 9", add="money=100")
        with self.db.tx() as c:
            c.execute("INSERT INTO role_permissions(level,role_id) VALUES('MINISTER','77')")
        self.econ2.roles = [discord.Role(77)]
        i = self.call(self.bank, "approvals", self.econ2)
        self.assertIn("Approval #1", i.text())
        self.assertEqual(list(i.view().handlers), ["Approve", "Revoke"])
        self.assertIn("Approved request #1", self.press(i.view(), "Approve", self.econ2).text())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT approved_by FROM approval_requests WHERE id=1").fetchone()[0], "222")

    # ------------------------------------------------------------- bulk transfers
    def bulk_setup(self):
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": 100000000 + 500000000, "coal": 10000000}   # members own $1M; alliance owns the rest

    def send_bulk(self, text, user=None, reason="weekly pay", auto_confirm=True):
        att = discord.Attachment("pay.csv", text.encode())
        return self.call(self.bulk, "send", user or self.admin, att, reason, auto_confirm=auto_confirm)

    def batch(self, bid=1):
        with self.db.read() as c:
            b = dict(c.execute("SELECT * FROM bulk_batches WHERE id=?", (bid,)).fetchone())
            items = [dict(r) for r in c.execute("SELECT * FROM bulk_items WHERE batch_id=? ORDER BY row_no", (bid,))]
        return b, items

    def test_bulk_good_file_pays_each_row_once_from_alliance_funds(self):
        self.bulk_setup()
        i = self.send_bulk("nation,money,coal,note\n2,1000,0,Pay\nGamma,0,500,Coal\n")
        self.assertEqual(self.pnw.withdraw_calls, 2)
        b, items = self.batch()
        self.assertEqual(b["status"], "COMPLETED")
        self.assertTrue(all(it["status"] == "COMPLETED" and it["tx_id"] for it in items))
        self.assertIn("Bulk transfer #1", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000)       # members untouched
            self.assertEqual(c.execute("SELECT COUNT(*) FROM transactions WHERE funding_source='ALLIANCE' "
                                       "AND status='COMPLETED'").fetchone()[0], 2)
        # same file again within 24h is refused: nothing is paid twice
        i = self.send_bulk("nation,money,coal,note\n2,1000,0,Pay\nGamma,0,500,Coal\n")
        self.assertIn("already sent as batch #1", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 2)
        # the results screen has item details
        view = self.call(self.bulk, "status", self.admin, 1).view()
        self.assertIn("Row 2", self.press(view, "Item results", self.admin).text())   # file line numbers

    def test_bulk_validation_blocks_everything_before_sending(self):
        self.bulk_setup()
        bad = ("nation,money,coal\n2,100,0\n2,5,0\n"        # same nation twice
               "99,100,0\n"                                    # not in the alliance
               "nobody-here,100,0\n"                            # unknown name
               "3,-5,0\n"                                       # negative
               "3,abc,0\n")                                     # malformed
        i = self.send_bulk(bad)
        t = i.text()
        for word in ("blocked", "already on row", "not in the alliance", "couldn't find", "negative", "not a valid number"):
            self.assertIn(word, t, word)
        self.assertEqual(self.pnw.withdraw_calls, 0)
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM bulk_batches").fetchone()[0], 0)

    def test_bulk_cannot_spend_member_money(self):
        self.fund_alice({"money": 1000000})                     # bank holds exactly what members own
        i = self.send_bulk("nation,money\n2,100\n")
        self.assertIn("short", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)

    def test_bulk_cancel_sends_nothing_and_staff_only(self):
        self.bulk_setup()
        self.send_bulk("nation,money\n2,100\n", auto_confirm=False)
        self.assertEqual(self.pnw.withdraw_calls, 0)
        i = self.send_bulk("nation,money\n2,100\n", user=self.alice)
        self.assertIn("permission", i.text())
        i = self.call(self.bulk, "template", self.admin)
        self.assertTrue(any(m.get("file") for m in i.sent))

    def test_bulk_pnw_rejections_are_reported_per_row(self):
        self.bulk_setup()
        self.pnw.withdraw_mode = "reject"
        self.send_bulk("nation,money\n2,100\n3,200\n")
        b, items = self.batch()
        self.assertEqual(b["status"], "PARTIAL")
        self.assertEqual([it["status"] for it in items], ["FAILED", "FAILED"])
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000)

    def test_bulk_halts_on_emergency_lock_and_resume_never_repeats(self):
        self.bulk_setup()
        with self.db.tx() as c:
            L.set_emergency_lock(c, True, "drill", "1")
        self.send_bulk("nation,money\n2,100\n3,200\n")
        b, items = self.batch()
        self.assertEqual(b["status"], "HALTED")
        self.assertEqual([it["status"] for it in items], ["PENDING", "PENDING"])
        self.assertEqual(self.pnw.withdraw_calls, 0)
        with self.db.tx() as c:
            L.set_emergency_lock(c, False, "", "1")
        self.call(self.bulk, "resume", self.admin, 1)
        self.assertEqual(self.pnw.withdraw_calls, 2)
        self.assertEqual(self.batch()[0]["status"], "COMPLETED")
        i = self.call(self.bulk, "resume", self.admin, 1)               # nothing left: nothing repeated
        self.assertEqual(self.pnw.withdraw_calls, 2)
        self.assertIn("nothing is waiting", i.text())

    def test_bulk_unconfirmed_transfer_stops_the_batch_and_never_resends(self):
        self.bulk_setup()
        self.pnw.withdraw_mode = "timeout_lost"
        self.send_bulk("nation,money\n2,100\n3,200\n")
        b, items = self.batch()
        self.assertEqual([it["status"] for it in items], ["UNCERTAIN", "PENDING"])
        self.assertEqual(self.pnw.withdraw_calls, 1)
        self.pnw.withdraw_mode = "ok"
        # While one transfer is unconfirmed the whole bank refuses new withdrawals (safety), so resume waits.
        self.call(self.bulk, "resume", self.admin, 1)
        self.assertEqual(self.pnw.withdraw_calls, 1)
        self.assertEqual(self.batch()[0]["status"], "HALTED")
        # Staff confirm (against PnW's records) that row 1 never left; then the rest can continue.
        with self.db.tx() as c:
            tx_id = c.execute("SELECT tx_id FROM bulk_items WHERE row_no=2").fetchone()[0]
            self.assertTrue(L.fail_tx(c, tx_id, "checked PnW: nothing sent"))
        self.call(self.bulk, "resume", self.admin, 1)
        b, items = self.batch()
        self.assertEqual(self.pnw.withdraw_calls, 2)                      # only the waiting row was sent
        self.assertEqual([it["status"] for it in items], ["FAILED", "COMPLETED"])
        self.assertEqual(b["status"], "PARTIAL")

    def test_bulk_large_batches_need_a_second_approver_and_limits_apply(self):
        self.bulk_setup()
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "approval_threshold_value", "500", "t")
            c.execute("INSERT INTO role_permissions(level,role_id) VALUES('MINISTER','77')")
        self.econ2.roles = [discord.Role(77)]
        csv = "nation,money\n2,100000\n"
        i = self.send_bulk(csv)
        self.assertIn("Approval request", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.call(self.bank, "approve", self.econ2, 1)
        self.send_bulk(csv)
        self.assertEqual(self.pnw.withdraw_calls, 1)
        # a per-transfer limit for the staff role blocks the oversize row up front
        with self.db.tx() as c:
            c.execute("INSERT INTO role_permissions(level,role_id) VALUES('BANKER','88')")
            c.execute("INSERT INTO limits(scope,scope_id,per_tx_cents) VALUES('ROLE','88',5000)")
        banker = discord.User(333, "banker")
        banker.roles = [discord.Role(88)]
        i = self.send_bulk("nation,money\n3,500\n", user=banker)
        self.assertIn("per-transfer limit", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 1)

    # ---------------------------------------------------------------- tax extras
    def add_old_tax(self, nation, days_ago, money):
        import datetime as dt
        d = (dt.date.today() - dt.timedelta(days=days_ago)).isoformat()
        rid = 8000 + days_ago
        with self.db.tx() as c:
            c.execute("INSERT INTO pnw_records(id,record_date,sender_id,sender_type,receiver_id,receiver_type,note,tax_id,"
                      "amounts_json,raw_json,raw_sha256,direction,classification,status,first_seen_at) "
                      "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                      (rid, d, nation, 1, 900, 2, "", 3, '{"money": %d}' % (money * 100), "{}", "x", "IN", "TAX", "NO_CREDIT", "x"))
            c.execute("INSERT INTO tax_records(pnw_record_id,nation_id,tax_id,record_date,amounts_json,recorded_at) VALUES(?,?,?,?,?,?)",
                      (rid, nation, 3, d, '{"money": %d}' % (money * 100), d))

    def test_tax_period_filters_and_member_export(self):
        self.fund_alice({"money": 1000})
        self.pnw.recs.append(rec(980, 1, {"money": 500}, "tax", tax_id=3))
        self.run_async(self.svc.scanner.scan())
        self.add_old_tax(1, 100, 4000)
        P = lambda v: discord.app_commands.Choice(v, v)       # noqa: E731
        t30 = self.call(self.tax, "dashboard", self.admin, P("30d")).text()
        tall = self.call(self.tax, "dashboard", self.admin, P("all")).text()
        self.assertIn("$500.00", t30)
        self.assertNotIn("$4,500.00", t30)
        self.assertIn("$4,500.00", tall)
        view = self.call(self.tax, "dashboard", self.admin).view()
        self.assertEqual(list(view.handlers), ["7 days", "30 days", "90 days", "All time", "Export", "Chart"])
        self.assertIn("$4,500.00", self.press(view, "All time", self.admin).text())
        self.assertIn("$500.00", self.call(self.tax, "report", self.admin, "alice", P("7d")).text())
        i = self.call(self.tax, "export", self.admin, P("all"))
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO([m["file"] for m in i.sent if m.get("file")][0].fp.getvalue()))
        self.assertEqual(wb.sheetnames, ["By member", "Records", "Info"])
        rows = list(wb["By member"].iter_rows(values_only=True))
        self.assertEqual(rows[1][0], 1)
        self.assertEqual(rows[1][4], 2)                                  # two tax records for nation 1
        self.assertEqual(len(list(wb["Records"].iter_rows())), 3)
        # members can never see any of it
        for name in ("dashboard", "profile", "brackets", "export"):
            args = ("alice",) if name == "profile" else ()
            self.assertIn("permission", self.call(self.tax, name, self.alice, *args).text(), name)

    def test_tax_brackets_profile_and_exemptions(self):
        self.fund_alice({"money": 1000})
        self.pnw.recs.append(rec(981, 1, {"money": 500}, "tax", tax_id=3))
        self.run_async(self.svc.scanner.scan())
        i = self.call(self.tax, "brackets", self.admin)
        for word in ("Core", "25%", "Newbies", "10%", "1 nation(s) paid"):
            self.assertIn(word, i.text())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM tax_brackets").fetchone()[0], 2)
        # if PnW is down, the saved brackets are still shown with a warning
        from tunbank.pnw import PnWUncertain

        async def down():
            raise PnWUncertain("down")
        self.pnw.fetch_tax_brackets = down
        i = self.call(self.tax, "brackets", self.admin)
        self.assertIn("Could not load brackets", i.text())
        self.assertIn("Core", i.text())
        # profile: bracket comes from the real tax record + saved rates
        i = self.call(self.tax, "profile", self.admin, "alice")
        for word in ("Tax profile", "#3", "Core", "cash 25%", "Not exempt", "$500.00"):
            self.assertIn(word, i.text())
        # exemptions: staff-only, tracked, never touches PnW or balances
        A = discord.app_commands.Choice
        add, rem, lst = A("add", "add"), A("remove", "remove"), A("list", "list")
        self.assertIn("permission", self.call(self.tax, "exemptions", self.alice, add, "alice", "x").text())
        i = self.call(self.tax, "exemptions", self.admin, add, "alice", "Founding member, tax-free", 30)
        self.assertIn("exemption added", i.text())
        self.assertIn("already has an active exemption", self.call(self.tax, "exemptions", self.admin, add, "alice", "again").text())
        self.assertIn("Exempt: Founding member", self.call(self.tax, "profile", self.admin, "alice").text())
        self.assertIn("exempt", self.call(self.tax, "dashboard", self.admin).text())
        self.assertIn("Founding member", self.call(self.tax, "exemptions", self.admin, lst).text())
        self.assertEqual(self.pnw.withdraw_calls, 0)
        # button path on the profile: end the exemption through the pop-up form
        view = self.call(self.tax, "profile", self.admin, "alice").view()
        i = self.press(view, "End exemption", self.admin)
        i.modal.children[0].value = "left the alliance council"
        self.run_async(i.modal.on_submit(FI(self.admin)))
        self.assertIn("Not exempt", self.call(self.tax, "profile", self.admin, "alice").text())
        self.assertIn("no active exemption", self.call(self.tax, "exemptions", self.admin, rem, "alice", "x").text())
        with self.db.read() as c:       # history is kept, never deleted
            self.assertEqual(c.execute("SELECT COUNT(*) FROM tax_exemptions").fetchone()[0], 1)

    def open_help(self, user):
        from tunbank import bot as botmod
        b = botmod.TunBankBot(self.settings, self.db)
        self.run_async(b.setup_hook())
        i = FI(user)
        self.run_async(b.tree.top["help"].callback(i))
        return b, i, i.view()

    # ---------------------------------------------------------------- grants
    def alliance_money(self, cash=500000000):
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": 100000000 + cash, "steel": 10000000}

    def test_grant_is_an_alliance_expenditure_with_a_full_record(self):
        self.alliance_money()
        self.assertIn("permission", self.call(self.grant, "send", self.alice, "2", "money=1m", "x").text())
        i = self.call(self.grant, "send", self.admin, "Beta", "money=1000000 steel=500", "Iron Dome project", project="Iron Dome")
        self.assertIn("Grant #1", i.text())
        self.assertIn("Completed", i.text())
        with self.db.read() as c:
            g = dict(c.execute("SELECT * FROM grants WHERE id=1").fetchone())
            tx = dict(c.execute("SELECT * FROM transactions WHERE id=?", (g["tx_id"],)).fetchone())
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000)        # no member money used
        self.assertEqual((g["status"], g["recipient_nation_id"], g["project"], g["requested_by"]), ("COMPLETED", 2, "Iron Dome", "900001"))
        self.assertEqual((tx["funding_source"], tx["status"]), ("ALLIANCE", "COMPLETED"))
        self.assertEqual(g["pnw_record_id"], tx["pnw_record_id"])
        self.assertGreater(g["value_cents"], 100000000)                                       # market value stored
        self.assertIn("Iron Dome", self.call(self.grant, "view", self.admin, 1).text())
        self.assertIn("Grant", self.call(self.grant, "list", self.admin).text())
        with self.db.tx() as c:
            with self.assertRaises(Exception):
                c.execute("UPDATE grants SET purpose='changed' WHERE id=1")
            with self.assertRaises(Exception):
                c.execute("DELETE FROM grants")

    def test_grant_needs_purpose_confirmation_funds_and_second_approver(self):
        self.alliance_money(cash=1000000)
        self.assertIn("needs a purpose", self.call(self.grant, "send", self.admin, "2", "money=1000", " ").text())
        self.call(self.grant, "send", self.admin, "2", "money=1000", "why", auto_confirm=False)
        self.assertEqual(self.pnw.withdraw_calls, 0)
        i = self.call(self.grant, "send", self.admin, "2", "money=999999999", "too big")        # more than the alliance owns
        self.assertIn("Failed", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "approval_threshold_value", "100", "t")
            c.execute("INSERT INTO role_permissions(level,role_id) VALUES('MINISTER','77')")
        self.econ2.roles = [discord.Role(77)]
        i = self.call(self.grant, "send", self.admin, "2", "money=5000", "Defence")
        self.assertIn("Approval request", i.text())
        self.call(self.bank, "approve", self.econ2, 1)
        i = self.call(self.grant, "send", self.admin, "2", "money=5000", "Defence")
        with self.db.read() as c:
            g = dict(c.execute("SELECT * FROM grants WHERE status='COMPLETED'").fetchone())
        self.assertEqual(g["approver"], "222")

    def test_offshore_without_configuration_explains_setup(self):
        i = self.call(self.bank, "offshore", self.admin)
        self.assertIn("OFFSHORE_ALLIANCE_ID", i.text())

    def test_help_shows_each_person_only_their_categories(self):
        b, i, hv = self.open_help(self.alice)
        self.assertIn("Welcome to TUN Bank", i.text())
        self.assertIn("Member", i.text())
        cats = [k for k in hv.handlers if k in ("Start here", "My account", "Banking", "Bulk transfers", "Tax",
                                               "Audit & security", "Charts", "Configuration")]
        self.assertEqual(cats, ["Start here", "My account"])
        t = self.press(hv, "My account", self.alice)
        t = t.edited["embed"]
        body = t.description
        self.assertIn("/bank withdrawself", body)
        self.assertNotIn("/bankset", body)
        self.assertNotIn("/bank reserve", body)
        # somebody else cannot flip my help pages
        self.assertIn("belongs to someone else", self.press(hv, "My account", self.admin).text())

    def test_help_admin_sees_every_category_with_pages_and_badges(self):
        b, i, hv = self.open_help(self.admin)
        cats = [k for k in hv.handlers if k in ("Start here", "My account", "Banking", "Bulk transfers", "Tax",
                                               "Audit & security", "Charts", "Configuration")]
        self.assertEqual(len(cats), 8)
        cfg = self.press(hv, "Configuration", self.admin).edited["embed"]
        self.assertIn("/bankset addbanker", cfg.description)
        self.assertIn("Admin", cfg.description)
        self.assertEqual(cfg.footer_text, "Configuration · page 1 of 3 · TUN Bank")
        nxt = self.press(hv, "Next page", self.admin).edited["embed"]
        self.assertIn("page 2 of 3", nxt.footer_text)
        self.assertIn("/bankset setrole", self.press(hv, "Next page", self.admin).edited["embed"].description)
        self.press(hv, "Banking", self.admin)
        pages = []
        for _ in range(3):
            pages.append(hv.card().description)
            if hv.index < len(hv.pages["banking"]) - 1:
                self.press(hv, "Next page", self.admin)
        self.assertTrue(any("/bank withdraw" in t for t in pages))
        self.assertTrue(any("/bank offshore" in t for t in pages))
        self.assertTrue(any("/grant send" in t for t in pages))

    def test_help_can_never_go_stale(self):
        """Every real command must be in the help catalog, and the catalog must not list ghosts."""
        from tunbank import bot as botmod
        from tunbank import cmds_help as H
        b = botmod.TunBankBot(self.settings, self.db)
        self.run_async(b.setup_hook())
        real = {p for p, _ in H.iter_commands(b.tree)}
        self.assertEqual(sorted(real - set(H.CATALOG)), [], "commands missing from /help (add them to CATALOG in cmds_help.py)")
        self.assertEqual(sorted(set(H.CATALOG) - real), [], "CATALOG lists commands that no longer exist")
        for path, (cat, lvl) in H.CATALOG.items():
            self.assertIn(cat, H.CATEGORIES, path)
            self.assertIn(lvl, H.ORDER, path)
        # and an Admin really sees all of them across the pages
        shown = set()
        for key in H.CATEGORIES:
            if key == "start":
                continue
            for card in H.category_cards(b.tree, {"ADMIN", "FLAG:bank_view_alliance_holdings", "FLAG:bank_view_tax"}, key):
                shown |= set(re.findall(r"\*\*/([a-z]+(?: [a-z]+)?)\*\*", card.description))
        self.assertEqual(sorted(real - shown), [])
        # descriptions come from the commands themselves
        card = H.category_cards(b.tree, {"ADMIN", "FLAG:bank_view_alliance_holdings"}, "bulk")[0]
        self.assertIn(b.tree.cmds["bulk"].commands["send"].description.replace("ECON: ", "")[:30].lower(), card.description.lower())

    def test_bot_wiring_registers_without_clashes(self):
        from tunbank import bot as botmod
        b = botmod.TunBankBot(self.settings, self.db)
        self.run_async(b.setup_hook())
        self.assertEqual(sorted(b.tree.cmds), ["audit", "bank", "bankset", "bulk", "chart", "deposit", "grant", "ledger", "loan", "nation", "tax"])


class OffshoreCmd(Cmd):
    """Same bot, configured with a separate offshore alliance."""
    OFF = 777

    def make_settings(self):
        return Settings("t", "k", "b", "k", ALLIANCE, Path(self.tmp.name), {900001},
                        offshore=BankAccess("offshore", self.OFF, "offkey", "offkey", "offbot"),
                        main_bot_key="mainbot", main_bot_api_key="mainapi")

    def setUp(self):
        super().setUp()
        self.pnw.off_holdings = {}

    def seed(self, main_cash=10**10):
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": main_cash, "coal": 10**8}

    def off_rows(self):
        with self.db.read() as c:
            return [dict(r) for r in c.execute("SELECT * FROM offshore_transfers ORDER BY id")]

    def test_funds_move_to_offshore_without_touching_any_member_balance(self):
        self.seed()
        before = self.snapshot_balances()
        i = self.call(self.bank, "offshore", self.admin, "money=5000000", "collect at the offshore")
        self.assertIn("Offshore transfer #1", i.text())
        self.assertIn("Member balances", i.text())
        self.assertEqual(self.pnw.last_bank, "main")                           # sent FROM the main bank...
        row = self.off_rows()[0]
        self.assertEqual((row["status"], row["mode"], row["direction"]), ("COMPLETED", "AUTO", "TO_OFFSHORE"))
        self.assertTrue(row["pnw_record_id"])                                  # ...tied to a real PnW record
        sent = self.pnw.recs[-1]
        self.assertEqual((sent["receiver_id"], sent["receiver_type"], sent["sender_id"]), (self.OFF, 2, ALLIANCE))
        self.assertIn("TUN-OFF1", sent["note"])
        self.assertEqual(self.snapshot_balances(), before)                     # ledger identical
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM ledger_entries WHERE entry_type!='DEPOSIT'").fetchone()[0], 0)
        self.run_async(self.svc.scanner.scan())                                # the scanner sees it again: no duplicate
        self.assertEqual(len(self.off_rows()), 1)

    def snapshot_balances(self):
        with self.db.read() as c:
            return [tuple(r) for r in c.execute("SELECT nation_id,bucket,resource,amount FROM balances ORDER BY 1,2,3")]

    def test_member_withdrawals_are_paid_from_the_offshore_bank(self):
        self.seed()
        self.call(self.bank, "offshore", self.admin, "money=5000000", "fund the offshore")
        self.pnw.off_holdings = {"money": 5000000}                              # PnW now shows it there
        i = self.call(self.bank, "withdrawself", self.alice, "money=40000")
        self.assertIn("Withdrawal completed", i.text())
        self.assertEqual(self.pnw.last_bank, "offshore")
        self.assertEqual(self.pnw.recs[-1]["sender_id"] if self.pnw.recs[-1]["sender_id"] == self.OFF else self.pnw.off_recs[-1]["sender_id"], self.OFF)
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000 - 4000000)
        # not enough physically at the offshore -> refused BEFORE anything is sent, balance intact
        calls = self.pnw.withdraw_calls
        self.pnw.off_holdings = {"money": 100}
        i = self.call(self.bank, "withdrawself", self.alice, "money=40000")
        self.assertIn("can't be processed right now", i.text())               # the member learns nothing about the bank
        self.assertNotIn("physically", i.text())
        self.assertNotIn("offshore", i.text().lower())
        self.assertEqual(self.pnw.withdraw_calls, calls)
        with self.db.read() as c:                                               # ...but ECON is told exactly why
            logged = [r["payload_json"] for r in c.execute("SELECT payload_json FROM alert_log")]
        self.assertTrue(any("does not physically hold enough" in x and "/bank offshore" in x for x in logged))

    def test_offshore_is_limited_to_what_the_main_bank_holds_and_the_keep_rule(self):
        self.seed(main_cash=100000000)
        i = self.call(self.bank, "offshore", self.admin, "money=999999999", "too much")
        self.assertIn("doesn't hold enough", i.text())
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "offshore_keep_in_main", "money=900000", "t")
        i = self.call(self.bank, "offshore", self.admin, "money=600000", "keeps 900k at main")
        self.assertIn("kept in the main bank by policy", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.assertEqual(self.off_rows(), [])

    def test_a_prepared_manual_transfer_stays_pending_and_the_database_refuses_to_fake_completion(self):
        self.seed()
        object.__setattr__(self.settings, "main_bot_key", "")
        i = self.call(self.bank, "offshore", self.admin, "money=1000", "manual")
        self.assertEqual(self.pnw.withdraw_calls, 0)
        row = self.off_rows()[0]
        self.assertEqual((row["status"], row["pnw_record_id"], row["completed_at"]), ("PLANNED", None, None))
        self.assertIn("Waiting for the real PnW transfer", self.call(self.bank, "offshore", self.admin).text())
        self.run_async(self.svc.scanner.scan())                                   # scanning with no matching record changes nothing
        self.assertEqual(self.off_rows()[0]["status"], "PLANNED")
        for sql in ("UPDATE offshore_transfers SET status='COMPLETED' WHERE id=1",
                    "UPDATE offshore_transfers SET status='COMPLETED', completed_at='x' WHERE id=1"):
            with self.assertRaises(Exception):
                with self.db.tx() as c:
                    c.execute(sql)
        self.assertEqual(self.off_rows()[0]["status"], "PLANNED")
        # a record with the tag but the WRONG direction/amount does not complete it either
        wrong = rec(7400, ALLIANCE, {"money": 5.0}, "TUN-OFF1", receiver=self.OFF, sender_type=2, receiver_type=2)
        self.pnw.recs.append(wrong)
        self.pnw.off_recs.append(wrong)
        self.run_async(self.svc.scanner.scan())
        self.assertEqual([r["status"] for r in self.off_rows()][0], "PLANNED")

    def test_receiver_type_is_a_setting_and_pnws_rejection_is_shown_verbatim(self):
        self.seed()
        self.assertEqual(self.settings.alliance_receiver_type, 2)
        object.__setattr__(self.settings, "alliance_receiver_type", 9)
        seen = {}
        orig = self.pnw.bank_withdraw

        async def strict(receiver, amounts, note, receiver_type=1, bank=None):
            seen["type"] = receiver_type
            raise PnWRejected('Variable "$receiver_type" got invalid value 9; expected 1 or 2')
        self.pnw.bank_withdraw = strict
        i = self.call(self.bank, "offshore", self.admin, "money=1000", "type test")
        self.assertEqual(seen["type"], 9)                                          # the configured value is what was sent
        self.assertIn("got invalid value 9", i.text())                              # PnW's own words
        self.assertIn("ALLIANCE_RECEIVER_TYPE", i.text())
        self.assertEqual(self.off_rows()[0]["status"], "FAILED")                    # never "completed"
        self.pnw.bank_withdraw = orig

    def test_who_may_offshore_is_configurable_default_staff_only(self):
        self.seed()
        self.assertIn("permission", self.call(self.bank, "offshore", self.alice, "money=1", "x").text())
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "offshore_access", "ADMIN", "t")
            c.execute("INSERT INTO bankers(discord_id,added_by,added_at) VALUES('222','x','x')")
        self.assertIn("permission", self.call(self.bank, "offshore", self.econ2, "money=1", "x").text())      # banker, but ADMIN-only now
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "offshore_access", "MEMBERS", "t")
        i = self.call(self.bank, "offshore", self.alice, "money=1", "x", auto_confirm=False)
        self.assertNotIn("permission", i.text())
        self.assertIn("Cancelled", i.text())

    def test_manual_mode_prepares_the_transfer_and_completes_when_pnw_shows_it(self):
        self.seed()
        self.svc.settings  # noqa: B018
        object.__setattr__(self.settings, "main_bot_key", "")                    # no main-alliance bot key
        object.__setattr__(self.settings, "main_bot_api_key", "")
        before = self.snapshot_balances()
        i = self.call(self.bank, "offshore", self.admin, "money=2000000", "manual run")
        self.assertIn("TUN-OFF1", i.text())
        self.assertIn(f"alliance #{self.OFF}", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 0)                             # the bot sent nothing
        self.assertEqual(self.off_rows()[0]["status"], "PLANNED")
        # ECON does it in-game with the tag
        real = rec(7001, ALLIANCE, {"money": 2000000.0}, "manual run TUN-OFF1", receiver=self.OFF, sender_type=2, receiver_type=2)
        self.pnw.recs.append(real)
        self.pnw.off_recs.append(real)
        r = self.run_async(self.svc.scanner.scan())
        self.assertEqual([o.kind for o in r.outcomes if o.record_id == 7001], ["OFFSHORE"])
        row = self.off_rows()[0]
        self.assertEqual((row["status"], row["pnw_record_id"]), ("COMPLETED", 7001))
        self.assertEqual(self.snapshot_balances(), before)

    def test_plans_can_be_cancelled_but_not_after_they_were_sent(self):
        self.seed()
        object.__setattr__(self.settings, "main_bot_key", "")
        i = self.call(self.bank, "offshore", self.admin, "money=1000", "plan")
        self.assertEqual(self.press(i.view(), "Cancel this plan", self.admin).text().count("cancelled"), 1)
        self.assertEqual(self.off_rows()[0]["status"], "CANCELLED")

    def test_in_game_transfers_between_our_banks_are_recorded_not_credited(self):
        self.seed()
        before = self.snapshot_balances()
        real = rec(7100, ALLIANCE, {"money": 3000000.0, "coal": 500.0}, "", receiver=self.OFF, sender_type=2, receiver_type=2)
        self.pnw.recs.append(real)
        self.pnw.off_recs.append(real)
        r = self.run_async(self.svc.scanner.scan())
        self.assertEqual([o.kind for o in r.outcomes if o.record_id == 7100], ["OFFSHORE"])
        self.assertEqual([x["mode"] for x in self.off_rows()], ["OBSERVED"])
        self.assertEqual(self.snapshot_balances(), before)
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM integrity_events WHERE kind='EXTERNAL_OUTFLOW'").fetchone()[0], 0)
        # a record that carries a plan's tag but the wrong amounts is flagged
        self.call(self.bank, "offshore", self.admin, "money=1000", "will be wrong")
        bad = rec(7101, ALLIANCE, {"money": 9999.0}, "TUN-OFF2", receiver=self.OFF, sender_type=2, receiver_type=2)
        self.pnw.recs.append(bad)
        self.pnw.off_recs.append(bad)
        self.run_async(self.svc.scanner.scan())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM integrity_events WHERE kind='OFFSHORE_MISMATCH' AND status='OPEN'").fetchone()[0], 1)

    def test_unclear_answers_never_resend_and_complete_when_the_record_appears(self):
        self.seed()
        self.pnw.withdraw_mode = "timeout_lost"
        i = self.call(self.bank, "offshore", self.admin, "money=1000", "flaky")
        self.assertIn("Uncertain", i.text())
        self.assertEqual(self.pnw.withdraw_calls, 1)
        self.assertEqual(self.off_rows()[0]["status"], "UNCERTAIN")
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM integrity_events WHERE kind='OFFSHORE_UNCERTAIN' AND status='OPEN'").fetchone()[0], 1)
        real = rec(7200, ALLIANCE, {"money": 1000.0}, "flaky TUN-OFF1", receiver=self.OFF, sender_type=2, receiver_type=2)
        self.pnw.recs.append(real)
        self.pnw.off_recs.append(real)
        self.run_async(self.svc.scanner.scan())
        self.assertEqual(self.off_rows()[0]["status"], "COMPLETED")
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM integrity_events WHERE kind='OFFSHORE_UNCERTAIN' AND status='OPEN'").fetchone()[0], 0)
        self.assertEqual(self.pnw.withdraw_calls, 1)
        # PnW did send it but the answer was lost: found immediately via its tag
        self.pnw.withdraw_mode = "timeout_sent"
        self.call(self.bank, "offshore", self.admin, "money=2000", "lost answer")
        self.assertEqual(self.off_rows()[1]["status"], "COMPLETED")
        self.assertEqual(self.pnw.withdraw_calls, 2)

    def test_pnw_refusal_and_emergency_lock(self):
        self.seed()
        self.pnw.withdraw_mode = "reject"
        i = self.call(self.bank, "offshore", self.admin, "money=1000", "refused")
        self.assertIn("Failed", i.text())
        self.assertEqual(self.off_rows()[0]["status"], "FAILED")
        with self.db.tx() as c:
            L.set_emergency_lock(c, True, "drill", "1")
        calls = self.pnw.withdraw_calls
        i = self.call(self.bank, "offshore", self.admin, "money=1000", "locked")
        self.assertIn("EMERGENCY LOCK", i.text())
        self.assertEqual(self.pnw.withdraw_calls, calls)

    def test_a_deposit_made_straight_to_the_offshore_goes_to_review_not_to_a_balance(self):
        self.seed()
        direct = rec(7300, 1, {"money": 5000.0}, "", receiver=self.OFF, sender_type=1, receiver_type=2)
        self.pnw.off_recs.append(direct)
        r = self.run_async(self.svc.scanner.scan())
        self.assertEqual([o.kind for o in r.outcomes if o.record_id == 7300], ["REVIEW"])
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000)

    def test_vault_shows_both_banks_and_reconciliation_uses_the_combined_total(self):
        self.seed(main_cash=100000000)
        self.pnw.off_holdings = {"money": 50000000}
        t = self.call(self.bank, "holdings", self.admin).text()
        for word in ("Main bank", "Offshore bank", "combined"):
            self.assertIn(word, t)
        with self.db.tx() as c:
            res = R.run_checks(c, holdings={"money": 150000000}, snapshot_id=None, triggered_by="t")
        self.assertNotIn("LEDGER_EXCEEDS_BANK", {f["kind"] for f in res["findings"]})
        # ownership vs physical: members own $1M; if BOTH banks together hold less, that is critical
        i = self.call(self.bank, "reconcile", self.admin)
        self.pnw.holdings, self.pnw.off_holdings = {"money": 10}, {"money": 10}
        i = self.call(self.bank, "reconcile", self.admin)
        with self.db.read() as c:
            self.assertTrue(L.integrity_state(c)["emergency_lock"])

    def test_offshore_status_screen_and_export(self):
        self.seed()
        self.call(self.bank, "offshore", self.admin, "money=1000", "first")
        i = self.call(self.bank, "offshore", self.admin)
        for word in ("Offshore", "Main bank", "Offshore bank", "Automatic", "Staff (Banker and above)"):
            self.assertIn(word, i.text())
        self.assertIn("Move funds…", i.view().handlers)
        self.assertTrue(any(m.get("file") for m in self.call(self.bank, "records", self.admin, discord.app_commands.Choice("offshore", "offshore")).sent))

    def test_config_refuses_to_guess_credentials(self):
        import os
        from tunbank.config import ConfigError, load_settings
        base = {"DISCORD_TOKEN": "t", "OWNER_DISCORD_IDS": "1", "ALLIANCE_ID": "10", "PNW_API_KEY": "k", "PNW_BOT_KEY": "b", "DATA_DIR": self.tmp.name}
        saved = dict(os.environ)
        try:
            os.environ.clear()
            os.environ.update(base)
            self.assertIsNone(load_settings().offshore)                         # no offshore: unchanged behaviour
            os.environ["OFFSHORE_ALLIANCE_ID"] = "10"
            with self.assertRaises(ConfigError) as e:
                load_settings()
            self.assertIn("DIFFERENT", str(e.exception))
            os.environ["OFFSHORE_ALLIANCE_ID"] = "20"
            with self.assertRaises(ConfigError) as e:
                load_settings()
            self.assertIn("OFFSHORE_API_KEY", str(e.exception))
            os.environ.update({"OFFSHORE_API_KEY": "ok", "OFFSHORE_BOT_KEY": "ob"})
            s = load_settings()
            self.assertEqual((s.offshore.alliance_id, s.payout.name, s.main.bot_key), (20, "offshore", None))   # main can't send
            self.assertEqual(s.bank_ids, {10, 20})
            os.environ.update({"MAIN_BOT_KEY": "mb", "MAIN_BOT_API_KEY": "ma"})
            self.assertEqual((load_settings().main.bot_key, load_settings().main.api_key), ("mb", "ma"))
        finally:
            os.environ.clear()
            os.environ.update(saved)


class FakeChannel:
    def __init__(self):
        self.sent = []

    async def send(self, embed=None, **kw):
        self.sent.append(embed)

    def text(self):
        return "\n".join(f"{e.title} {e.description} " + " ".join(f"{n} {v}" for n, v, _ in e.fields) for e in self.sent)


class FakeBotChannels:
    def __init__(self):
        self.chans = {}

    def get_channel(self, i):
        return self.chans.get(i)


class SecurityCmd(Cmd):
    """Who may see what, nation linking, guided deposits, tax-turn alerts."""

    def setUp(self):
        super().setUp()
        self.bot = FakeBotChannels()
        self.econ_log, self.tax_chan = FakeChannel(), FakeChannel()
        self.bot.chans = {111000: self.econ_log, 222000: self.tax_chan}
        self.svc.alerts.bot = self.bot
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "econ_log_channel_id", "111000", "t")
            cfg_set(c, "tax_alert_channel_id", "222000", "t")
            cfg_set(c, "tax_alert_settle_seconds", "0", "t")
            for lvl, role in (("AUDITOR", "66"), ("MINISTER", "77"), ("BANKER", "88")):
                c.execute("INSERT INTO role_permissions(level,role_id) VALUES(?,?)", (lvl, role))
        def person(uid, name, role):
            u = discord.User(uid, name)
            u.roles = [discord.Role(role)] if role else []
            return u
        self.auditor, self.minister, self.banker = person(401, "aud", 66), person(402, "min", 77), person(403, "bnk", 88)
        self.finance = person(404, "fin", 55)               # a role that only has a confidential permission

    def give(self, flag, role=55):
        self.call(self.bankset, "setaccess", self.admin, discord.app_commands.Choice(flag, flag), discord.Role(role))

    HOLD = "bank_view_alliance_holdings"

    # ------------------------------------------------------------- treasury confidentiality
    def treasury_calls(self, user):
        out = []
        out.append(self.call(self.bank, "holdings", user))
        out.append(self.call(self.bank, "offshore", user))
        out.append(self.call(self.chart, "vault", user))
        out.append(self.call(self.chart, "members", user))
        out.append(self.call(self.bank, "records", user, discord.app_commands.Choice("vault", "vault")))
        out.append(self.call(self.bank, "records", user, discord.app_commands.Choice("offshore", "offshore")))
        return out

    def test_members_cannot_see_alliance_holdings_any_way(self):
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": 987654321}
        for i in self.treasury_calls(self.alice):
            t = i.text()
            self.assertIn("confidential", t.lower(), t)
            self.assertNotIn("9,876,543", t)
            self.assertFalse(any(m.get("file") for m in i.sent))
        # a member's own screens never contain bank figures
        t = self.call(self.bank, "dashboard", self.alice).text() + self.call(self.bank, "balance", self.alice).text()
        self.assertNotIn("9,876,543", t)
        self.assertNotIn("ALLIANCE-OWNED", t)

    def test_staff_without_the_permission_cannot_see_holdings_even_ministers(self):
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": 987654321}
        for who in (self.auditor, self.banker, self.minister):
            for i in self.treasury_calls(who):
                self.assertIn("confidential", i.text().lower(), i.text())
                self.assertNotIn("9,876,543", i.text())

    def test_authorised_financial_role_and_admin_can_see_holdings(self):
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": 987654321}
        self.give(self.HOLD)
        t = self.call(self.bank, "holdings", self.finance).text()
        self.assertIn("9,876,543", t)
        self.assertIn("ALLIANCE-OWNED", t)
        self.call(self.bank, "reconcile", self.admin)
        self.assertTrue(any(m.get("file") for m in self.call(self.chart, "vault", self.finance).sent))
        self.assertTrue(any(m.get("file") for m in self.call(self.bank, "records", self.finance, discord.app_commands.Choice("vault", "vault")).sent))
        self.assertIn("9,876,543", self.call(self.bank, "holdings", self.admin).text())          # Admins always
        # the permission is separate from staff levels and from the tax permission
        self.assertIn("confidential", self.call(self.tax, "dashboard", self.finance).text().lower())
        self.assertIn("confidential", self.call(self.bank, "holdings", self.minister).text().lower())
        # removing the role takes it away again
        self.call(self.bankset, "setaccess", self.admin, discord.app_commands.Choice(self.HOLD, self.HOLD), discord.Role(55), True)
        self.assertIn("confidential", self.call(self.bank, "holdings", self.finance).text().lower())
        # only Admins configure access
        self.assertIn("permission", self.call(self.bankset, "setaccess", self.minister, discord.app_commands.Choice(self.HOLD, self.HOLD), discord.Role(77)).text())

    def test_tax_information_has_its_own_permission(self):
        self.fund_alice({"money": 1000})
        for name in ("dashboard", "brackets", "turns", "export"):
            self.assertIn("confidential", self.call(self.tax, name, self.minister).text().lower(), name)
        self.assertIn("confidential", self.call(self.chart, "tax", self.minister).text().lower())
        self.assertIn("confidential", self.call(self.bank, "records", self.minister, discord.app_commands.Choice("tax", "tax")).text().lower())
        self.give("bank_view_tax", 77)
        self.assertIn("Tax dashboard", self.call(self.tax, "dashboard", self.minister).text())
        self.assertIn("confidential", self.call(self.bank, "holdings", self.minister).text().lower())      # tax access != treasury access

    def test_treasury_numbers_do_not_leak_through_error_messages(self):
        self.alliance_money = None
        self.fund_alice({"money": 1000000})
        self.pnw.holdings = {"money": 100000000 + 500000}                    # alliance owns exactly $5,000
        # a Banker (no treasury permission) tries to spend more than the alliance owns
        i = self.call(self.bank, "withdraw", self.banker, discord.app_commands.Choice("ALLIANCE", "ALLIANCE"), "2", "money=999999", "x")
        self.assertIn("not sufficient", i.text())
        self.assertNotIn("5,000", i.text())
        i = self.send_bulk_as(self.banker, "nation,money\n2,999999\n")
        self.assertIn("not sufficient", i.text())
        self.assertNotIn("5,000.00", i.text())
        # an authorised role sees the real figures
        self.give(self.HOLD, 88)
        i = self.call(self.bank, "withdraw", self.banker, discord.app_commands.Choice("ALLIANCE", "ALLIANCE"), "2", "money=999999", "x")
        self.assertIn("$5,000.00", i.text())
        # the shared ECON log never carries bank figures from integrity alerts
        from tunbank.alerts import integrity_card
        card = integrity_card({"severity": "CRITICAL", "kind": "LEDGER_EXCEEDS_BANK", "event_id": 1, "message": "members hold more than the bank",
                               "details": {"shortfall": {"money": 123456789}, "bank": {"money": 5}, "member_total": {"money": 999}}})
        self.assertNotIn("123456789", card.as_text())

    def send_bulk_as(self, user, text):
        return self.call(self.bulk, "send", user, discord.Attachment("p.csv", text.encode()), "why")

    # ------------------------------------------------------------------- who sees which account
    def test_econ_staff_see_other_members_accounts_but_members_only_their_own(self):
        self.fund_alice({"money": 1000000, "coal": 500})
        for who in (self.auditor, self.banker, self.minister):
            t = self.call(self.audit, "nation", who, "alice").text()
            self.assertIn("$1,000,000.00", t)                                  # another member's deposit, no treasury access needed
            self.assertIn("confidential", self.call(self.bank, "holdings", who).text().lower())
        # a plain member cannot look at anybody (including themselves through the staff command)
        for name, group, args in (("nation", self.audit, ("alice",)), ("transactions", self.audit, ()), ("stafflog", self.audit, ())):
            self.assertIn("permission", self.call(group, name, self.alice, *args).text(), name)
        self.assertIn("permission", self.call(self.chart, "nation", self.alice, "alice", discord.app_commands.Choice("Resource mix", "mix")).text())
        # the only account a member can open is their own
        self.call(self.nation, "link", discord.User(222, "bob"), 2)
        t = self.call(self.bank, "dashboard", discord.User(222, "bob")).text()
        self.assertNotIn("$1,000,000.00", t)
        self.assertIn("Your TUN Bank account", t)                                   # bob sees only his own (empty) account

    # ----------------------------------------------------------------------- linking
    def test_admin_links_a_nation_to_a_member_and_everything_is_logged(self):
        bob = discord.User(222, "bob")
        i = self.call(self.bank, "linknation", self.admin, bob, "Beta")
        self.assertIn("Nation linked", i.text())
        with self.db.read() as c:
            m = dict(L.get_member(c, 2))
            hist = [dict(r) for r in c.execute("SELECT * FROM nation_link_history")]
            audit = [dict(r) for r in c.execute("SELECT * FROM audit_log WHERE action='NATION_LINKED_BY_ADMIN'")]
        self.assertEqual((m["discord_id"], m["discord_name"]), ("222", "bob"))
        self.assertEqual((hist[0]["action"], hist[0]["actor"], hist[0]["verified"]), ("LINKED", "900001", 1))   # PnW's discord field = bob
        self.assertEqual(len(audit), 1)
        # a Minister may create a brand-new link; a plain member may not
        self.assertIn("Nation linked", self.call(self.bank, "linknation", self.minister, discord.User(556, "dave"), "3").text())
        self.assertIn("permission", self.call(self.bank, "linknation", self.alice, discord.User(557, "eve"), "1").text())
        # nations outside the alliance are refused
        self.assertIn("not in our alliance", self.call(self.bank, "linknation", self.admin, discord.User(558, "fay"), "55").text())

    def test_duplicate_links_are_detected_and_need_an_admin_with_force(self):
        bob, mallory = discord.User(222, "bob"), discord.User(666, "mallory")
        self.call(self.bank, "linknation", self.admin, bob, "2")
        # someone else tries to take bob's nation
        i = self.call(self.bank, "linknation", self.minister, mallory, "2")
        self.assertIn("Link conflict", i.text())
        self.assertIn("already linked to <@222>", i.text())
        # a Minister cannot override, even with force
        i = self.call(self.bank, "linknation", self.minister, mallory, "2", force=True)
        self.assertIn("Link conflict", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_member(c, 2)["discord_id"], "222")
        # the same Discord account cannot silently take a second nation either
        i = self.call(self.bank, "linknation", self.admin, bob, "3")
        self.assertIn("already linked to nation [#2]", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_member(c, 2)["discord_id"], "222")
        # an Admin can replace it deliberately, with both people shown and the change logged
        i = self.call(self.bank, "linknation", self.admin, mallory, "2", force=True)
        self.assertIn("Nation linked", i.text())
        with self.db.read() as c:
            self.assertEqual(L.get_member(c, 2)["discord_id"], "666")
            acts = [r["action"] for r in c.execute("SELECT action FROM nation_link_history ORDER BY id")]
        self.assertEqual(acts, ["LINKED", "RELINKED"])
        # cancelling on the confirmation changes nothing
        self.call(self.bank, "linknation", self.admin, bob, "2", force=True, auto_confirm=False)
        with self.db.read() as c:
            self.assertEqual(L.get_member(c, 2)["discord_id"], "666")

    # -------------------------------------------------------------- guided deposit
    def test_deposit_command_never_creates_money_only_the_real_pnw_record_does(self):
        self.run_async(self.svc.scanner.scan())
        self.call(self.nation, "link", self.alice, 1)
        before = self.snapshot_ledger()
        i = self.call(self.bank, "deposit", self.alice, "money=5m coal=2000")
        t = i.text()
        for word in ("Deposit exactly", "leave the **note empty**", "Want the bot to do this for you?", "without a key you chose to give it"):
            self.assertIn(word, t)
        self.assertEqual(self.snapshot_ledger(), before)                         # no balance, no ledger line
        self.assertEqual(self.pnw.withdraw_calls, 0)
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT status FROM deposit_intents").fetchone()[0], "WAITING")
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE"), {})
        # pressing "check now" before depositing still creates nothing
        self.press(i.view(), "Check my deposit now", self.alice)
        self.assertEqual(self.snapshot_ledger(), before)
        # the withdraw route cannot spend a planned deposit
        self.assertEqual(self.pnw.withdraw_calls, 0)
        self.assertIn("don't have enough", self.call(self.bank, "withdrawself", self.alice, "money=1m").text())
        # a bad plan is refused
        self.assertIn("couldn't read", self.call(self.bank, "deposit", self.alice, "lots").text())
        # the REAL record arrives: now (and only now) the balance exists and the plan is matched
        self.pnw.recs.append(rec(9100, 1, {"money": 5000000.0, "coal": 2000.0}))
        self.run_async(self.svc.scanner.scan())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 500000000)
            row = dict(c.execute("SELECT * FROM deposit_intents").fetchone())
        self.assertEqual((row["status"], row["pnw_record_id"]), ("MATCHED", 9100))

    def snapshot_ledger(self):
        with self.db.read() as c:
            return (c.execute("SELECT COUNT(*) FROM ledger_entries").fetchone()[0],
                    [tuple(r) for r in c.execute("SELECT nation_id,bucket,resource,amount FROM balances")])

    # ---------------------------------------------------------------- tax-turn alerts
    def tax_rec(self, rid, nation, money, when, tax_id=3, note=""):
        r = rec(rid, nation, {"money": money, "coal": 10.0}, note, tax_id=tax_id)
        r["date"] = when
        return r

    def test_tax_alert_reports_only_the_turn_total_and_no_member_list(self):
        self.run_async(self.svc.scanner.scan())                                   # baseline
        self.pnw.recs += [self.tax_rec(9201, 1, 1000.0, "2026-10-02 14:00:02"),
                          self.tax_rec(9202, 2, 2000.0, "2026-10-02 14:00:03"),
                          self.tax_rec(9203, 3, 3000.0, "2026-10-02 14:00:04")]
        r = self.run_async(self.svc.scanner.scan())
        from tunbank.ui import post_outcomes
        self.run_async(post_outcomes(self.svc, r.outcomes))
        self.assertEqual(len(self.tax_chan.sent), 1)                              # ONE message for the whole turn
        t = self.tax_chan.text()
        for word in ("Tax Collection — Turn Complete", "02 Oct 2026 — 14:00 UTC", "$6,000.00", "Cash", "Resources", "Current Market Value"):
            self.assertIn(word, t)
        for leak in ("Alpha", "Beta", "Gamma", "[#1]", "[#2]", "[#3]", "<@"):
            self.assertNotIn(leak, t)
        self.assertNotIn("Tax collection recorded", self.econ_log.text())         # no per-record posts anywhere
        self.assertEqual(len(self.econ_log.sent), 0)
        # scanning again never repeats the alert
        r = self.run_async(self.svc.scanner.scan())
        self.run_async(post_outcomes(self.svc, r.outcomes))
        self.assertEqual(len(self.tax_chan.sent), 1)
        # the next turn is a new alert; the detail is available to authorised staff only
        self.pnw.recs.append(self.tax_rec(9204, 1, 500.0, "2026-10-02 16:00:01"))
        r = self.run_async(self.svc.scanner.scan())
        self.run_async(post_outcomes(self.svc, r.outcomes))
        self.assertEqual(len(self.tax_chan.sent), 2)
        self.assertIn("16:00 UTC", self.tax_chan.sent[1].description)
        self.give("bank_view_tax", 77)
        detail = self.call(self.tax, "report", self.minister, "alpha").text() + self.call(self.tax, "turns", self.minister).text()
        self.assertIn("$1,500.00", detail)

    def test_tax_is_recognised_from_pnw_tax_records_not_from_the_word_tax(self):
        self.run_async(self.svc.scanner.scan())
        self.pnw.recs += [rec(9301, 1, {"money": 500.0}, "tax payment for the alliance"),                 # a member's note, no PnW tax id
                          self.tax_rec(9302, 2, 700.0, "2026-10-02 18:00:01", note="")]                    # real PnW tax record, empty note
        r = self.run_async(self.svc.scanner.scan())
        kinds = {o.record_id: o.kind for o in r.outcomes}
        self.assertEqual(kinds[9301], "CREDIT")                                                            # an ordinary deposit
        self.assertEqual(kinds[9302], "TAX")
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM tax_records").fetchone()[0], 1)
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 50000)
            self.assertEqual(L.get_balances(c, 2, "AVAILABLE"), {})                                        # tax never credits a member

    def test_old_tax_history_found_on_the_first_scan_never_triggers_alerts(self):
        self.pnw.recs.append(self.tax_rec(9401, 1, 1000.0, "2026-09-30 10:00:01"))
        r = self.run_async(self.svc.scanner.scan())                                                        # first scan = baseline
        from tunbank.ui import post_outcomes
        self.run_async(post_outcomes(self.svc, r.outcomes))
        self.assertEqual(len(self.tax_chan.sent), 0)

    # ------------------------------------------------------- configuration / security audit log
    def audit_rows(self, **where):
        with self.db.read() as c:
            q = "SELECT * FROM config_audit"
            if where:
                q += " WHERE " + " AND ".join(f"{k}=?" for k in where)
            return [dict(r) for r in c.execute(q + " ORDER BY id", tuple(where.values()))]

    def audit_channel(self):
        self.cfg_chan = FakeChannel()
        self.bot.chans[333000] = self.cfg_chan
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "config_audit_channel_id", "333000", "t")
        self.run_async(self.svc.alerts.flush_config_audit())          # post the setup changes so tests start clean
        self.cfg_chan.sent.clear()

    def test_a_setting_change_is_recorded_with_who_what_before_after_and_posted_privately(self):
        self.audit_channel()
        C = discord.app_commands.Choice
        self.call(self.bankset, "config", self.admin, "offshore_access", "ADMIN")
        self.run_async(self.svc.alerts.flush_config_audit())
        self.cfg_chan.sent.clear()
        self.call(self.bankset, "config", self.admin, "offshore_access", "STAFF")
        rows = [r for r in self.audit_rows(setting="offshore_access")]
        last = rows[-1]
        self.assertEqual((last["actor_id"], last["previous"], last["new"], last["action"], last["category"]),
                         ("900001", "ADMIN", "STAFF", "/bankset config", "FINANCIAL"))
        self.assertRegex(last["ts"], r"^\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ$")
        self.run_async(self.svc.alerts.flush_config_audit())
        self.assertEqual(len(self.cfg_chan.sent), 1)
        t = self.cfg_chan.text()
        for needle in ("Financial Setting Changed", "<@900001>", "`900001`", "offshore_access", "`ADMIN`", "`STAFF`", "/bankset config", "UTC"):
            self.assertIn(needle, t)
        self.assertEqual(self.audit_rows(setting="offshore_access")[-1]["posted_at"] is not None, True)
        self.assertNotIn("Configuration Changed", self.econ_log.text())            # kept apart from normal transaction logs
        # setting the same value again changes nothing, so nothing is logged
        n = len(self.audit_rows())
        self.call(self.bankset, "config", self.admin, "offshore_access", "STAFF")
        self.assertEqual(len(self.audit_rows()), n)
        # only Admins can read the history
        self.assertIn("permission", self.call(self.audit, "configlog", self.minister).text())
        self.assertIn("offshore_access", self.call(self.audit, "configlog", self.admin, "offshore").text())
        self.assertTrue(any(m.get("file") for m in self.call(self.bank, "records", self.admin, C("configaudit", "configaudit")).sent))
        self.assertIn("Admin", self.call(self.bank, "records", self.minister, C("configaudit", "configaudit")).text())

    def test_every_important_admin_action_is_logged_never_silent(self):
        self.fund_alice({"money": 1000000})
        C, R = discord.app_commands.Choice, discord.Role
        bob = discord.User(222, "bob")
        steps = [
            ("role permission", lambda: self.call(self.bankset, "setrole", self.admin, C("BANKER", "BANKER"), R(91)), "PERMISSION", "role_permission:BANKER"),
            ("role removal", lambda: self.call(self.bankset, "setrole", self.admin, C("BANKER", "BANKER"), R(91), True), "PERMISSION", "role_permission:BANKER"),
            ("confidential access", lambda: self.give("bank_view_tax", 92), "PERMISSION", "confidential_access:bank_view_tax"),
            ("add banker", lambda: self.call(self.bankset, "addbanker", self.admin, discord.User(77, "x")), "PERMISSION", "banker_access"),
            ("remove banker", lambda: self.call(self.bankset, "removebanker", self.admin, discord.User(77, "x")), "PERMISSION", "banker_access"),
            ("global limit", lambda: self.call(self.bankset, "settransferlimit", self.admin, "500m"), "LIMIT", "limit:GLOBAL"),
            ("daily limit", lambda: self.call(self.bankset, "setdailylimit", self.admin, "1b"), "LIMIT", "limit:GLOBAL"),
            ("role limit", lambda: self.call(self.bankset, "setrolelimit", self.admin, R(88), "10m"), "LIMIT", "limit:ROLE"),
            ("nation limit", lambda: self.call(self.bankset, "setnationlimit", self.admin, "alice", "2m"), "LIMIT", "limit:NATION"),
            ("approval threshold", lambda: self.call(self.bankset, "requireapproval", self.admin, 5000), "FINANCIAL", "approval_threshold_value"),
            ("offshore keep rule", lambda: self.call(self.bankset, "config", self.admin, "offshore_keep_in_main", "money=1m"), "FINANCIAL", "offshore_keep_in_main"),
            ("grant level", lambda: self.call(self.bankset, "config", self.admin, "grant_min_level", "ADMIN"), "FINANCIAL", "grant_min_level"),
            ("price setting", lambda: self.call(self.bankset, "config", self.admin, "price_stale_seconds", "999"), "CONFIG", "price_stale_seconds"),
            ("log channel", lambda: self.call(self.bankset, "setlogchannel", self.admin, discord.TextChannel(5)), "CHANNEL", "econ_log_channel_id"),
            ("icon", lambda: self.call(self.bankset, "seticon", self.admin, C("Oil", "oil"), "🔥"), "ICON", "icon_oil"),
            ("freeze", lambda: self.call(self.bank, "freeze", self.admin, "alice", "checking"), "SECURITY", "account_frozen"),
            ("unfreeze", lambda: self.call(self.bank, "unfreeze", self.admin, "alice", "all clear"), "SECURITY", "account_frozen"),
            ("pause", lambda: self.call(self.bank, "lock", self.admin, "maintenance"), "SECURITY", "withdrawals_paused"),
            ("resume", lambda: self.call(self.bank, "unlock", self.admin, "done"), "SECURITY", "withdrawals_paused"),
            ("emergency on", lambda: self.call(self.ledger, "emergencylock", self.admin, C("on", "on"), "drill"), "SECURITY", "emergency_financial_lock"),
            ("emergency off", lambda: self.call(self.ledger, "emergencylock", self.admin, C("off", "off"), "drill over"), "SECURITY", "emergency_financial_lock"),
            ("adjustment", lambda: self.call(self.bank, "adjust", self.admin, "alice", "typo", "ticket 4", remove="money=1"), "ACCOUNTING", "balance_adjustment"),
            ("link", lambda: self.call(self.bank, "linknation", self.admin, bob, "Beta"), "LINK", "nation_link"),
            ("tax exemption", lambda: self.call(self.tax, "exemptions", self.admin, C("add", "add"), "alice", "founder"), "POLICY", "tax_exemption"),
            ("opening import", lambda: self.call(self.bankset, "importopening", self.admin, discord.Attachment("o.csv", b"nation_id,money\n3,5\n"), "legacy"), "IMPORT", "opening_balance_import"),
        ]
        for label, run, category, setting in steps:
            before = len(self.audit_rows())
            run()
            new = self.audit_rows()[before:]
            self.assertTrue(new, f"{label}: nothing was logged")
            hit = [r for r in new if r["setting"] == setting]
            self.assertTrue(hit, f"{label}: expected {setting}, got {[r['setting'] for r in new]}")
            self.assertEqual(hit[0]["category"], category, label)
            self.assertEqual(hit[0]["actor_id"], "900001", label)
            self.assertTrue(hit[0]["action"].startswith("/"), f"{label}: action {hit[0]['action']!r}")
        # manual classification of a PnW record
        self.pnw.recs.append(rec(9700, 77, {"money": 5}))
        self.run_async(self.svc.scanner.scan())
        before = len(self.audit_rows())
        self.call(self.bank, "review", self.admin, 9700, C("alliance", "alliance"), "", "was a donation")
        new = self.audit_rows()[before:]
        self.assertEqual([r["setting"] for r in new], ["record_classification"])
        self.assertEqual(new[0]["category"], "CLASSIFICATION")
        self.assertIn("PnW record #9700", new[0]["target"])

    def test_nothing_is_silent_when_the_channel_is_missing_or_down(self):
        # no audit channel configured: entries wait in the database and ECON is told once
        self.call(self.bankset, "config", self.admin, "grant_min_level", "ADMIN")
        self.call(self.bankset, "config", self.admin, "grant_min_level", "MINISTER")
        self.run_async(self.svc.alerts.flush_config_audit())
        self.run_async(self.svc.alerts.flush_config_audit())
        pending = [r for r in self.audit_rows() if r["posted_at"] is None]
        self.assertGreaterEqual(len(pending), 2)
        self.assertEqual(self.econ_log.text().count("Configuration audit channel not set"), 1)
        self.assertNotIn("`ADMIN`", self.econ_log.text())                         # the details never go to the normal log
        # the channel is set later: the backlog is posted, oldest first
        self.audit_channel_without_clearing()
        self.run_async(self.svc.alerts.flush_config_audit())
        self.assertEqual([r for r in self.audit_rows() if r["posted_at"] is None], [])
        self.assertIn("grant_min_level", self.cfg_chan.text())
        # a channel that errors keeps the entry queued for retry instead of dropping it
        class Broken(FakeChannel):
            async def send(self, embed=None, **kw):
                raise RuntimeError("discord is down")
        self.bot.chans[333000] = Broken()
        self.call(self.bankset, "config", self.admin, "grant_min_level", "BANKER")
        self.run_async(self.svc.alerts.flush_config_audit())
        self.assertEqual(len([r for r in self.audit_rows() if r["posted_at"] is None]), 1)
        self.bot.chans[333000] = self.cfg_chan
        self.run_async(self.svc.alerts.flush_config_audit())
        self.assertEqual([r for r in self.audit_rows() if r["posted_at"] is None], [])

    def audit_channel_without_clearing(self):
        self.cfg_chan = FakeChannel()
        self.bot.chans[333000] = self.cfg_chan
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "config_audit_channel_id", "333000", "t")

    def test_audit_entries_cannot_be_edited_or_deleted(self):
        self.call(self.bankset, "config", self.admin, "grant_min_level", "ADMIN")
        for sql in ("DELETE FROM config_audit", "UPDATE config_audit SET new='MINISTER'", "UPDATE config_audit SET actor_id='1'",
                    "UPDATE config_audit SET setting='x'"):
            with self.assertRaises(Exception):
                with self.db.tx() as c:
                    c.execute(sql)
        self.assertGreaterEqual(len(self.audit_rows()), 1)

    def test_changes_made_in_the_env_file_are_noticed_at_startup_without_leaking_keys(self):
        from tunbank import configaudit as CA
        with self.db.tx() as c:
            self.assertEqual(CA.check_env_changes(c, self.settings), 0)             # first start records the baseline
        object.__setattr__(self.settings, "alliance_receiver_type", 9)
        object.__setattr__(self.settings, "pnw_api_key", "SUPERSECRETKEY123")
        with self.db.tx() as c:
            self.assertEqual(CA.check_env_changes(c, self.settings), 2)
            self.assertEqual(CA.check_env_changes(c, self.settings), 0)             # not repeated
        rows = [r for r in self.audit_rows() if r["category"] == "ENVIRONMENT"]
        self.assertEqual({r["setting"] for r in rows}, {"alliance_receiver_type", "pnw_read_key"})
        everything = repr(rows)
        self.assertNotIn("SUPERSECRETKEY123", everything)                           # only a short fingerprint is kept
        self.assertTrue(all(r["actor_id"] == "system:startup" for r in rows))

    def test_existing_tax_history_is_grouped_without_announcing_it(self):
        from tunbank.records import backfill_tax_turns
        self.run_async(self.svc.scanner.scan())
        self.pnw.recs += [self.tax_rec(9601, 1, 100.0, "2026-10-02 10:00:01"), self.tax_rec(9602, 2, 200.0, "2026-10-02 10:00:02")]
        self.run_async(self.svc.scanner.scan())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT records FROM tax_turns").fetchone()[0], 2)
        # simulate the upgrade: an existing database with tax records but an empty tax_turns table
        raw = __import__("sqlite3").connect(str(self.db.path))
        raw.execute("DELETE FROM tax_turns")
        raw.commit()
        raw.close()
        with self.db.tx() as c:
            self.assertEqual(backfill_tax_turns(c), 1)
            self.assertEqual(backfill_tax_turns(c), 0)                       # only once
        self.run_async(self.svc.alerts.flush_tax_turns(self.svc.prices))
        self.assertEqual(len(self.tax_chan.sent), 0)                         # history is not re-announced
        with self.db.read() as c:
            row = c.execute("SELECT records, alerted_at FROM tax_turns").fetchone()
        self.assertEqual((row[0], row[1]), (2, "backfill"))

    def test_tax_alert_flags_missing_prices_instead_of_guessing(self):
        self.run_async(self.svc.scanner.scan())
        async def none():
            return None
        self.svc.prices.get = none
        self.pnw.recs.append(self.tax_rec(9501, 1, 1000.0, "2026-10-02 20:00:01"))
        r = self.run_async(self.svc.scanner.scan())
        self.run_async(self.svc.alerts.flush_tax_turns(self.svc.prices))
        self.assertIn("no price for: coal", self.tax_chan.text())                # honest about the missing price, never a silent zero


class DepositCmd(Cmd):
    """Members deposit from Discord using their OWN API key. Nothing is credited without the real PnW record."""
    KEY = "alicekey1234567890"
    BOT = "b"

    def make_settings(self):
        from cryptography.fernet import Fernet
        return Settings("t", "k", self.BOT, "k", ALLIANCE, Path(self.tmp.name), {900001}, credential_key=Fernet.generate_key().decode())

    def setUp(self):
        super().setUp()
        self.pnw.key_owner = {self.KEY: 1, "bobkey1234567890": 2}
        self.bot = FakeBotChannels()
        self.econ_log = FakeChannel()
        self.bot.chans = {111000: self.econ_log}
        self.svc.alerts.bot = self.bot
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "econ_log_channel_id", "111000", "t")
        CR.install_log_redaction(self.settings)

    def ledger_state(self):
        with self.db.read() as c:
            return (c.execute("SELECT COUNT(*) FROM ledger_entries").fetchone()[0],
                    [tuple(r) for r in c.execute("SELECT nation_id,bucket,resource,amount FROM balances")])

    def set_key(self, user=None, key=None):
        user = user or self.alice
        i = self.call(self.nation, "setkey", user)
        i.modal.children[0].value = key or self.KEY
        sub = FI(user)
        self.run_async(i.modal.on_submit(sub))
        return sub

    def ready(self):
        self.run_async(self.svc.scanner.scan())                       # baseline
        self.call(self.nation, "link", self.alice, 1)
        return self.set_key()

    def deps(self):
        with self.db.read() as c:
            return [dict(r) for r in c.execute("SELECT * FROM member_deposits ORDER BY id")]

    def test_key_is_stored_encrypted_checked_for_ownership_and_never_shown(self):
        self.run_async(self.svc.scanner.scan())
        self.call(self.nation, "link", self.alice, 1)
        self.assertIn("not linked your nation", self.call(self.nation, "setkey", discord.User(999, "stranger")).text().lower())   # unlinked people cannot
        sub = self.set_key()
        t = sub.text()
        self.assertIn("Your API key is saved", t)
        self.assertIn("Whitelisted access", t)
        self.assertIn(CR.hint(self.KEY), t)
        self.assertNotIn(self.KEY, t)                                           # never shown back
        with self.db.read() as c:
            row = dict(c.execute("SELECT * FROM member_credentials").fetchone())
            audit = repr([dict(r) for r in c.execute("SELECT * FROM config_audit")]) + repr([dict(r) for r in c.execute("SELECT * FROM audit_log")])
        self.assertEqual((row["nation_id"], row["discord_id"], row["verified"], row["key_hint"]), (1, "111", 1, CR.hint(self.KEY)))
        self.assertNotIn(self.KEY.encode(), bytes(row["key_enc"]))               # encrypted at rest
        self.assertNotIn(self.KEY, audit)                                        # not in any audit trail either
        self.assertIn("member_api_key", audit)
        with open(self.db.path, "rb") as fh:
            self.assertNotIn(self.KEY.encode(), fh.read())                                 # not anywhere in the database file
        # someone else's key (it belongs to nation 2) cannot be saved for nation 1
        sub = self.set_key(key="bobkey1234567890")
        self.assertIn("belongs to a different nation", sub.text())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT key_hint FROM member_credentials WHERE nation_id=1").fetchone()[0], CR.hint(self.KEY))
        # junk and keys PnW rejects are refused
        self.assertIn("doesn't look like", self.set_key(key="not a key!!").text())
        self.assertIn("didn't accept that key", self.set_key(key="unknownkey1234567").text())
        # removing deletes it
        self.assertIn("deleted", self.call(self.nation, "removekey", self.alice).text())
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM member_credentials").fetchone()[0], 0)

    def test_deposit_from_discord_credits_only_when_the_real_pnw_record_is_seen(self):
        self.ready()
        before = self.ledger_state()
        i = self.call(self.bank, "deposit", self.alice, "money=5m coal=2000")
        t = i.text()
        for word in ("Deposit Credited", "PnW record", "Deposit"):
            self.assertIn(word, t)
        call = self.pnw.deposit_calls[0]
        self.assertEqual((call["key"], call["bot"], call["note"]), (self.KEY, self.BOT, "TUN-DEP1"))   # the member's own key + the bot key
        self.assertEqual(call["amounts"], {"money": 500000000, "coal": 200000})
        self.assertNotIn(self.KEY, t)
        d = self.deps()[0]
        self.assertEqual((d["status"], d["nation_id"]), ("CREDITED", 1))
        self.assertIsNotNone(d["pnw_record_id"])                                  # the PnW transaction id is stored
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE"), {"money": 500000000, "coal": 200000})
            led = c.execute("SELECT pnw_record_id, entry_type FROM ledger_entries").fetchall()
        self.assertTrue(all(r[0] == d["pnw_record_id"] and r[1] == "DEPOSIT" for r in led))
        # scanning again, or the same record seen twice, never credits twice
        self.run_async(self.svc.scanner.scan())
        self.run_async(self.svc.scanner.scan())
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 500000000)
        # the same confirmation id can't start a second deposit
        self.run_async(self.svc.deposits.start(nation_id=1, discord_id=111, amounts={"money": 100}, idem="mdep-fixed"))
        res2 = self.run_async(self.svc.deposits.start(nation_id=1, discord_id=111, amounts={"money": 100}, idem="mdep-fixed"))
        self.assertEqual(len(self.pnw.deposit_calls), 2)                    # the repeat did not call PnW again
        self.assertIn("already handled", res2["message"])

    def test_confirmation_is_required_and_cancelling_sends_nothing(self):
        self.ready()
        i = self.call(self.bank, "deposit", self.alice, "money=1m", auto_confirm=False)
        self.assertIn("Cancelled", i.text())
        self.assertIsNone(self.pnw.deposit_calls)
        self.assertEqual(self.deps(), [])
        first = self.call(self.bank, "deposit", self.alice, "money=1m", auto_confirm=False).sent[0]["embed"]
        self.assertIn("Confirm deposit from your nation", first.title)

    def test_pnw_refusal_changes_nothing_and_explains_whitelisted_access(self):
        self.ready()
        self.pnw.deposit_mode = "reject"
        before = self.ledger_state()
        t = self.call(self.bank, "deposit", self.alice, "money=1m").text()
        self.assertIn("Failed", t)
        self.assertIn("PnW refused the deposit, nothing was sent", t)
        self.assertIn("not authorized", t)                                        # PnW's own words
        self.assertIn("Whitelisted access", t)
        self.assertEqual(self.ledger_state(), before)
        self.assertEqual(self.deps()[0]["status"], "FAILED")
        self.assertNotIn(self.KEY, t)

    def test_unclear_answers_are_never_resent_and_credit_only_from_the_record(self):
        self.ready()
        self.pnw.deposit_mode = "timeout_lost"
        t = self.call(self.bank, "deposit", self.alice, "money=1m").text()
        self.assertIn("Uncertain", t)
        self.assertEqual(len(self.pnw.deposit_calls), 1)
        self.assertEqual(self.deps()[0]["status"], "UNCERTAIN")
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE"), {})
            self.assertEqual(c.execute("SELECT COUNT(*) FROM integrity_events WHERE kind='MEMBER_DEPOSIT_UNCERTAIN' AND status='OPEN'").fetchone()[0], 1)
        # PnW did move the money after all: the real record shows up later and credits it, exactly once
        real = rec(8800, 1, {"money": 1000000.0}, "TUN-DEP1")
        self.pnw.recs.append(real)
        self.run_async(self.svc.scanner.scan())
        self.assertEqual(self.deps()[0]["status"], "CREDITED")
        self.assertEqual(self.deps()[0]["pnw_record_id"], 8800)
        with self.db.read() as c:
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE")["money"], 100000000)
            self.assertEqual(c.execute("SELECT COUNT(*) FROM integrity_events WHERE kind='MEMBER_DEPOSIT_UNCERTAIN' AND status='OPEN'").fetchone()[0], 0)
        self.assertEqual(len(self.pnw.deposit_calls), 1)
        # answer lost but PnW did take it: found straight away via the tag
        self.pnw.deposit_mode = "timeout_sent"
        t = self.call(self.bank, "deposit", self.alice, "money=2m").text()
        self.assertIn("Credited", t)
        self.assertEqual(len(self.pnw.deposit_calls), 2)

    def test_a_key_that_acts_as_another_nation_is_disabled_and_nothing_is_credited_to_the_wrong_person(self):
        self.ready()
        self.pnw.key_owner[self.KEY] = 2               # PnW now says this key belongs to Beta (nation 2)
        t = self.call(self.bank, "deposit", self.alice, "money=1m").text()
        self.assertIn("does not belong to your nation", t)
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT disabled FROM member_credentials").fetchone()[0], 1)
            self.assertEqual(L.get_balances(c, 1, "AVAILABLE"), {})                  # nothing for Alpha
            self.assertEqual(L.get_balances(c, 2, "AVAILABLE")["money"], 100000000)   # the real money came from Beta's nation
        self.assertEqual(self.deps()[0]["status"], "FAILED")
        with self.db.read() as c:
            self.assertIn("MEMBER_KEY_WRONG_NATION", [r[0] for r in c.execute("SELECT kind FROM integrity_events")])
        # a disabled key is not used again: the member gets the manual steps instead
        t = self.call(self.bank, "deposit", self.alice, "money=1m").text()
        self.assertIn("Deposit exactly", t)
        self.assertEqual(len(self.pnw.deposit_calls), 1)

    def test_a_key_only_works_for_its_own_member_and_nation(self):
        self.ready()
        # an Admin moves the nation to someone else: the old owner's key is deleted, the new owner can't use it
        mallory = discord.User(666, "mallory")
        self.call(self.bank, "linknation", self.admin, mallory, "1", force=True)
        with self.db.read() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM member_credentials").fetchone()[0], 0)
        t = self.call(self.bank, "deposit", mallory, "money=1m").text()
        self.assertIn("Deposit exactly", t)                                       # manual steps, no direct deposit
        self.assertIsNone(self.pnw.deposit_calls)
        # the loader itself refuses a mismatched member even if a row exists
        with self.db.tx() as c:
            CR.save(c, self.svc.crypto, nation_id=1, discord_id="111", api_key=self.KEY, verified=True)
            self.assertIsNone(CR.load_for_member(c, self.svc.crypto, nation_id=1, discord_id="666"))
            self.assertIsNone(CR.load_for_member(c, self.svc.crypto, nation_id=2, discord_id="111"))

    def test_without_encryption_or_when_switched_off_the_guided_steps_are_used(self):
        self.ready()
        with self.db.tx() as c:
            from tunbank.config import cfg_set
            cfg_set(c, "member_deposit_enabled", "0", "t")
        t = self.call(self.bank, "deposit", self.alice, "money=1m").text()
        self.assertIn("Deposit exactly", t)
        self.assertIsNone(self.pnw.deposit_calls)
        self.assertEqual(self.deps(), [])

    def test_secrets_are_scrubbed_from_logs_and_messages(self):
        import io
        import logging
        self.ready()
        stream = io.StringIO()
        h = logging.StreamHandler(stream)
        h.addFilter(CR.RedactingFilter())
        lg = logging.getLogger("tunbank.test_secret")
        lg.addHandler(h)
        lg.propagate = False
        lg.warning("PnW said no for key %s at https://x/?api_key=%s", self.KEY, "abcdef123456")
        self.assertNotIn(self.KEY, stream.getvalue())
        self.assertNotIn("abcdef123456", stream.getvalue())
        self.assertIn("hidden key", stream.getvalue())
        self.assertNotIn(self.KEY, CR.redact(f"error for {self.KEY}"))

    def test_no_internal_balance_without_a_real_pnw_transaction_even_if_the_code_is_misused(self):
        self.ready()
        with self.assertRaises(Exception):
            with self.db.tx() as c:                       # claiming a deposit is CREDITED with no PnW record is refused by the database
                c.execute("INSERT INTO member_deposits(nation_id,discord_id,amounts_json,status,created_at,updated_at,idempotency_key) "
                          "VALUES(1,'111','{}','CREDITED','x','x','fake')")
        with self.assertRaises(Exception):
            with self.db.tx() as c:
                L.post_entries(c, [dict(group_id="g", nation_id=1, bucket="AVAILABLE", resource="money", delta=10**9,
                                        entry_type="DEPOSIT", pnw_record_id=555, actor="x")])
        self.assertEqual(self.ledger_state()[0], 0)


for _cls in (OffshoreCmd, SecurityCmd, DepositCmd):
    for _name in dir(Cmd):                  # only each class's own tests run; inherited ones run once in Cmd
        if _name.startswith("test_") and _name not in _cls.__dict__:
            setattr(_cls, _name, None)


if __name__ == "__main__":
    unittest.main()